from openpyxl import load_workbook
from pathlib import Path
import json, re, hashlib
from io import BytesIO
base=Path(r'D:\天气地图\地质资料')
out=Path(r'C:\Users\sigeryang\Documents\ChatGPT\天气项目\output\geology-research-vertical\data')
files={'回次分层采样':out/'source-converted.xlsx','分析结果':base/'测试分析结果.xlsx','柱状图模板':out/'template-converted.xlsx'}
def extract(path,label):
    raw=path.read_bytes()
    wbv=load_workbook(BytesIO(raw),data_only=False,read_only=False)
    wbc=load_workbook(BytesIO(raw),data_only=True,read_only=False)
    result={'file':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':[]}
    for ws in wbv.worksheets:
        wc=wbc[ws.title]
        cells=[]
        for row in ws.iter_rows():
            for c in row:
                v=c.value
                if v is None: continue
                cells.append({'cell':c.coordinate,'value':v,'cached':wc[c.coordinate].value if isinstance(v,str) and v.startswith('=') else None,
                              'type':c.data_type,'number_format':c.number_format})
        merged=[str(x) for x in ws.merged_cells.ranges]
        heights={str(i):d.height for i,d in ws.row_dimensions.items() if d.height is not None}
        widths={k:d.width for k,d in ws.column_dimensions.items() if d.width is not None}
        result['sheets'].append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'state':ws.sheet_state,
          'merged':merged,'row_heights':heights,'column_widths':widths,'freeze_panes':str(ws.freeze_panes) if ws.freeze_panes else None,
          'cells':cells,'images':len(ws._images),'charts':len(ws._charts)})
    (out/(label+'.json')).write_text(json.dumps(result,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
    return result
allres={k:extract(p,k) for k,p in files.items()}
(out/'all-workbooks-extract.json').write_text(json.dumps(allres,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
for k,res in allres.items():
 print('\n##',k,res['file'])
 for s in res['sheets']:
  print(s['name'],s['max_row'],s['max_column'],'nonempty',len(s['cells']),'merged',len(s['merged']),'charts',s['charts'],'images',s['images'])




