from openpyxl import load_workbook
from pathlib import Path
p=Path(r'output\geology-research-vertical\data\template-converted.xlsx')
w=load_workbook(p,data_only=False)
for sn in ['回次','分层','样品']:
 s=w[sn]
 print(sn,'heights',[(i,s.row_dimensions[i].height) for i in range(1,s.max_row+1) if s.row_dimensions[i].height is not None][:15], 'distinct',sorted(set(s.row_dimensions[i].height for i in range(1,s.max_row+1) if s.row_dimensions[i].height is not None)))
 print('print_area',s.print_area,'orientation',s.page_setup.orientation,'fit',s.page_setup.fitToHeight,s.page_setup.fitToWidth)
 print('hidden rows',[(i,s.row_dimensions[i].hidden) for i in range(1,s.max_row+1) if s.row_dimensions[i].hidden])
print('calculation',w.calculation)
