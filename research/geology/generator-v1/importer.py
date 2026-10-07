"""Import measured geology traverses into the generator v1 JSON contract.

Only the supplied workbook is read.  Geometry is independently recomputed from
slant length, slope and azimuth; spreadsheet cache columns are audit evidence.
"""

from __future__ import annotations

import hashlib
import math
import os
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


SCHEMA_VERSION = "1.0"
CANONICAL_COLUMNS = [
    "记录号", "导线段号", "层号", "起读数_m", "止读数_m", "斜距_m", "坡角_deg",
    "方位角_deg", "倾向_deg", "倾角_deg", "岩性名称", "岩性描述", "样品编号",
    "样品距测段起点_m", "原表平距_m", "原表高差_m", "原表累计高差_m",
    "原表累计北_m", "原表累计东_m", "实测真厚度_m",
]
PROJECT_PARAMETERS = ["项目名称", "剖面编号", "剖面方位角_deg", "长度单位", "角度单位"]


class GeometryInputError(ValueError):
    """Raised when workbook geometry cannot be represented without invention."""

    def __init__(self, issues: List[Dict[str, Any]]):
        self.issues = issues
        errors = [i for i in issues if i.get("severity") == "error"]
        message = "; ".join(i.get("message", "invalid geometry") for i in errors[:5])
        super().__init__(message or "invalid geometry input")


def _issue(severity: str, code: str, message: str, cells: Iterable[str] = ()) -> Dict[str, Any]:
    return {"severity": severity, "code": code, "message": message, "cells": list(cells)}


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any, cell: str, issues: List[Dict[str, Any]], *, required: bool = False) -> Optional[float]:
    if _is_blank(value):
        if required:
            issues.append(_issue("error", "MISSING_GEOMETRY", f"{cell} 缺少必需几何值", [cell]))
        return None
    if isinstance(value, bool):
        issues.append(_issue("error", "INVALID_NUMBER", f"{cell} 的 TRUE/FALSE 不是有效数字", [cell]))
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        issues.append(_issue("error", "INVALID_NUMBER", f"{cell} 不是有效数字", [cell]))
        return None
    if not math.isfinite(result):
        issues.append(_issue("error", "NONFINITE_NUMBER", f"{cell} 不允许 NaN 或无穷值", [cell]))
        return None
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell(col: int, row: int) -> str:
    letters = ""
    n = col + 1
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def _load_xlrd():
    try:
        import xlrd  # type: ignore
        return xlrd
    except ImportError:
        candidates = []
        if os.environ.get("GEOLOGY_XLRD_PATH"):
            candidates.append(Path(os.environ["GEOLOGY_XLRD_PATH"]))
        candidates.append(
            Path(__file__).resolve().parents[2]
            / ".openai" / "geology-review-20261005" / "updated-materials" / "read-deps"
        )
        for candidate in candidates:
            if candidate.exists():
                sys.path.insert(0, str(candidate))
                try:
                    import xlrd  # type: ignore
                    return xlrd
                except ImportError:
                    pass
        raise ImportError("读取 .xls 需要 xlrd；可用 GEOLOGY_XLRD_PATH 指定依赖目录")


def _old_xls_rows(path: Path) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    xlrd = _load_xlrd()
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    sheet = workbook.sheet_by_index(0)
    issues: List[Dict[str, Any]] = []
    title = _text(sheet.cell_value(0, 0)) if sheet.nrows else ""
    header_a = _text(sheet.cell_value(2, 0)) if sheet.nrows > 2 else ""
    header_e = _text(sheet.cell_value(2, 4)) if sheet.nrows > 2 and sheet.ncols > 4 else ""
    if "实测地层剖面登记表" not in title.replace("　", "").replace(" ", "") or "皮尺编号" not in header_a.replace("　", "").replace(" ", "") or "斜距" not in header_e.replace("　", "").replace(" ", ""):
        raise GeometryInputError([_issue("error", "UNRECOGNIZED_XLS", "首表不符合已知实测地层剖面登记表表头，拒绝按 A:V 盲读", ["A1", "A3", "E3"])])

    section_heading = _text(sheet.cell_value(1, 0)) if sheet.nrows > 1 else ""
    section_id = section_heading.split()[0] if section_heading else ""
    project_name = _text(sheet.cell_value(1, 10)) if sheet.nrows > 1 and sheet.ncols > 10 else ""
    meta = {"name": project_name, "section_id": section_id, "sheet": sheet.name}

    rows: List[Dict[str, Any]] = []
    previous_leg = ""
    previous_leg_cell = ""
    previous_layer = ""
    previous_layer_cell = ""
    # Data starts after the five-row heading. A valid candidate has at least one
    # core measurement value; this excludes totals and signature/footer rows.
    for r0 in range(5, sheet.nrows):
        values = [sheet.cell_value(r0, c) if c < sheet.ncols else "" for c in range(22)]
        if not any(not _is_blank(values[c]) for c in (1, 2, 3, 4, 5, 9, 12, 13)):
            continue
        row_num = r0 + 1
        leg_explicit = _text(values[0])
        layer_explicit = _text(values[1])
        leg_id = leg_explicit or previous_leg
        layer_id = layer_explicit or previous_layer
        leg_source = _cell(0, row_num) if leg_explicit else previous_leg_cell
        layer_source = _cell(1, row_num) if layer_explicit else previous_layer_cell
        if leg_explicit:
            previous_leg, previous_leg_cell = leg_explicit, _cell(0, row_num)
        if layer_explicit:
            previous_layer, previous_layer_cell = layer_explicit, _cell(1, row_num)
        rows.append({
            "record_id": "",
            "leg_id": leg_id,
            "leg_explicit": bool(leg_explicit),
            "layer_id": layer_id,
            "start_reading": values[2], "end_reading": values[3],
            "length": values[4], "slope": values[5], "azimuth": values[9],
            "dip_direction": values[12], "dip_angle": values[13],
            "lithology_name": "",
            "description": values[15], "sample_id": values[20], "sample_offset": None,
            "true_thickness": values[17],
            "cache": {"horizontal_m": values[6], "vertical_m": values[7],
                      "cumulative_z_m": values[8], "cumulative_north_m": values[10],
                      "cumulative_east_m": values[11]},
            "raw": { _cell(c, row_num): values[c] for c in range(22) if not _is_blank(values[c]) },
            "source_cells": {
                "leg_id": leg_source or _cell(0, row_num), "layer_id": layer_source or _cell(1, row_num),
                "start_reading_m": _cell(2, row_num), "end_reading_m": _cell(3, row_num),
                "length_m": _cell(4, row_num), "slope_deg": _cell(5, row_num),
                "azimuth_deg": _cell(9, row_num), "dip_direction_deg": _cell(12, row_num),
                "dip_angle_deg": _cell(13, row_num), "description": _cell(15, row_num),
                "sample_id": _cell(20, row_num),
                "true_thickness_m": _cell(17, row_num),
                "cache_horizontal_m": _cell(6, row_num), "cache_vertical_m": _cell(7, row_num),
                "cache_cumulative_z_m": _cell(8, row_num), "cache_cumulative_north_m": _cell(10, row_num),
                "cache_cumulative_east_m": _cell(11, row_num),
            },
            "row": row_num,
        })
    workbook.release_resources()
    return meta, rows, issues


def _canonical_xlsx_rows(path: Path) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError("读取 .xlsx 需要 openpyxl") from exc
    issues: List[Dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "项目" not in workbook.sheetnames or "测段" not in workbook.sheetnames:
        workbook.close()
        raise GeometryInputError([_issue("error", "MISSING_SHEET", "规范 XLSX 必须包含“项目”和“测段”工作表", [])])
    project_sheet = workbook["项目"]
    if _text(project_sheet["A1"].value) != "参数" or _text(project_sheet["B1"].value) != "值":
        workbook.close()
        raise GeometryInputError([_issue("error", "INVALID_PROJECT_HEADER", "项目表 A1/B1 必须为“参数”/“值”", ["项目!A1", "项目!B1"])])
    params: Dict[str, Any] = {}
    param_cells: Dict[str, str] = {}
    for row, cells in enumerate(project_sheet.iter_rows(min_row=2), start=2):
        key = _text(cells[0].value if len(cells) >= 1 else None)
        if key:
            params[key] = cells[1].value if len(cells) >= 2 else None
            param_cells[key] = f"项目!B{row}"
    missing_params = [p for p in PROJECT_PARAMETERS if p not in params]
    if missing_params:
        issues.append(_issue("error", "MISSING_PROJECT_PARAMETER", f"项目表缺少参数：{', '.join(missing_params)}", ["项目!A:A"]))
    if _text(params.get("长度单位")) != "m":
        issues.append(_issue("error", "INVALID_LENGTH_UNIT", "长度单位必须为 m", [param_cells.get("长度单位", "项目!B:B")]))
    if _text(params.get("角度单位")) != "deg":
        issues.append(_issue("error", "INVALID_ANGLE_UNIT", "角度单位必须为 deg", [param_cells.get("角度单位", "项目!B:B")]))
    explicit_axis = params.get("剖面方位角_deg")
    if not _is_blank(explicit_axis):
        axis = _number(explicit_axis, param_cells.get("剖面方位角_deg", "项目!B:B"), issues)
        params["剖面方位角_deg"] = axis

    segment_sheet = workbook["测段"]
    headers = [_text(cell.value) for cell in next(segment_sheet.iter_rows(min_row=1, max_row=1))]
    duplicates = sorted({h for h in headers if h and headers.count(h) > 1})
    missing_columns = [name for name in CANONICAL_COLUMNS if name not in headers]
    if duplicates:
        issues.append(_issue("error", "DUPLICATE_COLUMN", f"测段表存在重复列名：{', '.join(duplicates)}", ["测段!1:1"]))
    if missing_columns:
        issues.append(_issue("error", "MISSING_COLUMN", f"测段表缺少列：{', '.join(missing_columns)}", ["测段!1:1"]))
    column = {name: headers.index(name) + 1 for name in CANONICAL_COLUMNS if name in headers}

    rows: List[Dict[str, Any]] = []
    previous_leg = previous_layer = ""
    previous_leg_cell = previous_layer_cell = ""
    for row_num, cells in enumerate(segment_sheet.iter_rows(min_row=2), start=2):
        values = {name: cells[idx - 1].value for name, idx in column.items() if idx <= len(cells)}
        if not any(not _is_blank(value) for value in values.values()):
            continue
        leg_explicit = _text(values.get("导线段号"))
        layer_explicit = _text(values.get("层号"))
        leg_cell = f"测段!{_cell(column.get('导线段号', 1) - 1, row_num)}"
        layer_cell = f"测段!{_cell(column.get('层号', 1) - 1, row_num)}"
        leg_id = leg_explicit or previous_leg
        layer_id = layer_explicit or previous_layer
        leg_source = leg_cell if leg_explicit else previous_leg_cell
        layer_source = layer_cell if layer_explicit else previous_layer_cell
        if leg_explicit:
            previous_leg, previous_leg_cell = leg_explicit, leg_cell
        if layer_explicit:
            previous_layer, previous_layer_cell = layer_explicit, layer_cell

        def ref(name: str) -> str:
            return f"测段!{_cell(column.get(name, 1) - 1, row_num)}"

        raw = {ref(name): value for name, value in values.items() if not _is_blank(value)}
        rows.append({
            "record_id": _text(values.get("记录号")), "leg_id": leg_id,
            "leg_explicit": bool(leg_explicit), "layer_id": layer_id,
            "start_reading": values.get("起读数_m"), "end_reading": values.get("止读数_m"),
            "length": values.get("斜距_m"), "slope": values.get("坡角_deg"),
            "azimuth": values.get("方位角_deg"), "dip_direction": values.get("倾向_deg"),
            "dip_angle": values.get("倾角_deg"), "lithology_name": _text(values.get("岩性名称")),
            "description": values.get("岩性描述"), "sample_id": values.get("样品编号"),
            "sample_offset": values.get("样品距测段起点_m"),
            "true_thickness": values.get("实测真厚度_m"),
            "cache": {"horizontal_m": values.get("原表平距_m"), "vertical_m": values.get("原表高差_m"),
                      "cumulative_z_m": values.get("原表累计高差_m"), "cumulative_north_m": values.get("原表累计北_m"),
                      "cumulative_east_m": values.get("原表累计东_m")},
            "raw": raw,
            "source_cells": {
                "record_id": ref("记录号"), "leg_id": leg_source or leg_cell,
                "layer_id": layer_source or layer_cell, "start_reading_m": ref("起读数_m"),
                "end_reading_m": ref("止读数_m"), "length_m": ref("斜距_m"),
                "slope_deg": ref("坡角_deg"), "azimuth_deg": ref("方位角_deg"),
                "dip_direction_deg": ref("倾向_deg"), "dip_angle_deg": ref("倾角_deg"),
                "lithology_name": ref("岩性名称"), "description": ref("岩性描述"),
                "sample_id": ref("样品编号"), "sample_offset_m": ref("样品距测段起点_m"),
                "true_thickness_m": ref("实测真厚度_m"),
                "cache_horizontal_m": ref("原表平距_m"), "cache_vertical_m": ref("原表高差_m"),
                "cache_cumulative_z_m": ref("原表累计高差_m"), "cache_cumulative_north_m": ref("原表累计北_m"),
                "cache_cumulative_east_m": ref("原表累计东_m"),
            },
            "row": row_num,
        })
    workbook.close()
    meta = {"name": _text(params.get("项目名称")), "section_id": _text(params.get("剖面编号")),
            "sheet": "测段", "axis": params.get("剖面方位角_deg")}
    return meta, rows, issues


def _parse_leg(value: str) -> Optional[Tuple[str, str]]:
    match = re.fullmatch(r"\s*([^\s\-—–]+)\s*[\-—–]\s*([^\s\-—–]+)\s*", value)
    return (match.group(1), match.group(2)) if match else None


def _cache_warning(issues: List[Dict[str, Any]], cell: str, label: str, raw: Any,
                   computed: float, tolerance: float) -> None:
    if _is_blank(raw):
        return
    try:
        raw_number = float(raw)
    except (TypeError, ValueError):
        issues.append(_issue("warning", "INVALID_CACHE_VALUE", f"{cell} 的{label}缓存不是数字，已保留原值", [cell]))
        return
    if not math.isfinite(raw_number):
        issues.append(_issue("warning", "INVALID_CACHE_VALUE", f"{cell} 的{label}缓存不是有限数，已保留原值", [cell]))
        return
    residual = raw_number - computed
    if abs(residual) > tolerance:
        issues.append(_issue("warning", "CACHE_MISMATCH", f"{label}缓存 {raw_number!r} 与独立计算 {computed!r} 的残差为 {residual!r}", [cell]))


def load_workbook_data(path: Any, settings: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Read a legacy measured-section XLS or the canonical XLSX format."""
    source_path = Path(path).expanduser().resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    user_settings = dict(settings or {})
    tolerance = user_settings.get("cache_tolerance_m", 1e-6)
    try:
        tolerance = float(tolerance)
    except (TypeError, ValueError):
        tolerance = -1.0
    if not math.isfinite(tolerance) or tolerance < 0:
        raise GeometryInputError([_issue("error", "INVALID_TOLERANCE", "cache_tolerance_m 必须是非负有限数", [])])
    origin = user_settings.get("origin")
    if origin not in (None, {}, {"east_m": 0, "north_m": 0, "z_m": 0}, {"east_m": 0.0, "north_m": 0.0, "z_m": 0.0}):
        raise GeometryInputError([_issue("error", "INVALID_ORIGIN", "v1 使用局部相对原点 0/0/0", [])])

    suffix = source_path.suffix.lower()
    if suffix == ".xls":
        meta, input_rows, issues = _old_xls_rows(source_path)
        adapter = "legacy_xls_v1"
    elif suffix == ".xlsx":
        meta, input_rows, issues = _canonical_xlsx_rows(source_path)
        adapter = "canonical_xlsx_v1"
    else:
        raise GeometryInputError([_issue("error", "UNSUPPORTED_FORMAT", "仅支持 .xls 与 .xlsx", [])])

    if not input_rows:
        issues.append(_issue("error", "NO_RECORDS", "工作簿没有测段记录", []))

    records: List[Dict[str, Any]] = []
    east = north = z = horizontal_chainage = slant_chainage = 0.0
    previous_description = previous_lithology = previous_layer = ""
    previous_description_cell = previous_lithology_cell = ""
    for index, row in enumerate(input_rows):
        cells = row["source_cells"]
        leg_id, layer_id = _text(row["leg_id"]), _text(row["layer_id"])
        if not leg_id:
            issues.append(_issue("error", "MISSING_LEG_ID", "测量行缺少导线段号且无法继承", [cells["leg_id"]]))
        if not layer_id:
            issues.append(_issue("error", "MISSING_LAYER_ID", "测量行缺少层号且无法继承", [cells["layer_id"]]))
        length = _number(row["length"], cells["length_m"], issues, required=True)
        slope = _number(row["slope"], cells["slope_deg"], issues, required=True)
        azimuth = _number(row["azimuth"], cells["azimuth_deg"], issues, required=True)
        dip_direction = _number(row["dip_direction"], cells["dip_direction_deg"], issues) if not _is_blank(row["dip_direction"]) else None
        dip_angle = _number(row["dip_angle"], cells["dip_angle_deg"], issues) if not _is_blank(row["dip_angle"]) else None
        if length is not None and length <= 0:
            issues.append(_issue("error", "INVALID_LENGTH", "斜距必须大于 0", [cells["length_m"]]))
        if slope is not None and not -90 <= slope <= 90:
            issues.append(_issue("error", "INVALID_SLOPE", "坡角必须在 [-90, 90]", [cells["slope_deg"]]))
        if azimuth is not None and not 0 <= azimuth < 360:
            issues.append(_issue("error", "INVALID_AZIMUTH", "方位角必须在 [0, 360)", [cells["azimuth_deg"]]))
        if (dip_direction is None) != (dip_angle is None):
            issues.append(_issue("error", "INCOMPLETE_ATTITUDE", "倾向和倾角必须同时提供或同时留空", [cells["dip_direction_deg"], cells["dip_angle_deg"]]))
        if dip_direction is not None and not 0 <= dip_direction < 360:
            issues.append(_issue("error", "INVALID_DIP_DIRECTION", "倾向必须在 [0, 360)", [cells["dip_direction_deg"]]))
        if dip_angle is not None and not 0 <= dip_angle <= 90:
            issues.append(_issue("error", "INVALID_DIP_ANGLE", "倾角必须在 [0, 90]", [cells["dip_angle_deg"]]))
        thickness = None
        if not _is_blank(row["true_thickness"]):
            thickness = _number(row["true_thickness"], cells.get("true_thickness_m", ""), issues)
            if thickness is not None and thickness < 0:
                issues.append(_issue("error", "INVALID_TRUE_THICKNESS", "实测真厚度不得小于 0", [cells.get("true_thickness_m", "")]))
        if length is None or slope is None or azimuth is None or length <= 0 or not -90 <= slope <= 90 or not 0 <= azimuth < 360:
            continue

        start_reading_raw, end_reading_raw = row["start_reading"], row["end_reading"]
        if _is_blank(start_reading_raw) != _is_blank(end_reading_raw):
            present_cell = cells["end_reading_m"] if _is_blank(start_reading_raw) else cells["start_reading_m"]
            issues.append(_issue("warning", "INCOMPLETE_READING_PAIR",
                                 "起读数和止读数只提供了一项，已保留原值且不推断另一项", [present_cell]))
        elif not _is_blank(start_reading_raw):
            reading_cells = [cells["start_reading_m"], cells["end_reading_m"]]
            if isinstance(start_reading_raw, bool) or isinstance(end_reading_raw, bool):
                issues.append(_issue("warning", "INVALID_READING_CACHE", "起止读数必须是有限数字，已保留原值", reading_cells))
            else:
                try:
                    start_reading, end_reading = float(start_reading_raw), float(end_reading_raw)
                    if not math.isfinite(start_reading) or not math.isfinite(end_reading):
                        raise ValueError
                except (TypeError, ValueError):
                    issues.append(_issue("warning", "INVALID_READING_CACHE", "起止读数必须是有限数字，已保留原值", reading_cells))
                else:
                    residual = (end_reading - start_reading) - length
                    if abs(residual) > tolerance:
                        issues.append(_issue("warning", "READING_LENGTH_MISMATCH",
                                             f"止读数减起读数与斜距的残差为 {residual!r}，保留原值并采用斜距计算",
                                             reading_cells + [cells["length_m"]]))

        slope_rad, azimuth_rad = math.radians(slope), math.radians(azimuth)
        horizontal = length * math.cos(slope_rad)
        vertical = length * math.sin(slope_rad)
        delta_east = horizontal * math.sin(azimuth_rad)
        delta_north = horizontal * math.cos(azimuth_rad)
        east += delta_east
        north += delta_north
        z += vertical
        horizontal_chainage += horizontal
        slant_chainage += length

        description = _text(row["description"])
        lithology = _text(row["lithology_name"])
        if adapter == "legacy_xls_v1" and description:
            lithology = description.split("。", 1)[0].strip()
        description_source = cells["description"]
        lithology_source = cells.get("lithology_name", cells["description"])
        same_layer = layer_id == previous_layer
        if not description and same_layer and previous_description:
            description, description_source = previous_description, previous_description_cell
        if not lithology and same_layer and previous_lithology:
            lithology, lithology_source = previous_lithology, previous_lithology_cell
        cells = dict(cells)
        cells["description"] = description_source
        cells["lithology_name"] = lithology_source
        # Establish independent forward-inheritance state for every new layer as
        # well as explicit changes within a layer. Explicit values are never
        # replaced by the other field's presence or absence.
        if not same_layer:
            previous_description = previous_lithology = ""
            previous_description_cell = previous_lithology_cell = ""
        if description:
            previous_description, previous_description_cell = description, description_source
        if lithology:
            previous_lithology, previous_lithology_cell = lithology, lithology_source
        previous_layer = layer_id

        record_id = _text(row["record_id"]) or f"R{index + 1:04d}"
        if any(existing["id"] == record_id for existing in records):
            issues.append(_issue("error", "DUPLICATE_RECORD_ID", f"记录号“{record_id}”重复", [cells.get("record_id", "")]))
        record = {
            "id": record_id, "leg_id": leg_id, "layer_id": layer_id,
            "length_m": length, "slope_deg": slope, "azimuth_deg": azimuth,
            "dip_direction_deg": dip_direction, "dip_angle_deg": dip_angle,
            "lithology_name": lithology, "description": description,
            "start_node": len(records), "end_node": len(records) + 1,
            "raw": row["raw"], "source_cells": cells,
            "computed": {"horizontal_m": horizontal, "vertical_m": vertical,
                         "east_delta_m": delta_east, "north_delta_m": delta_north},
            "true_thickness_m": thickness,
        }
        records.append(record)
        cache = row["cache"]
        _cache_warning(issues, cells["cache_horizontal_m"], "平距", cache["horizontal_m"], horizontal, tolerance)
        _cache_warning(issues, cells["cache_vertical_m"], "高差", cache["vertical_m"], vertical, tolerance)
        _cache_warning(issues, cells["cache_cumulative_z_m"], "累计高差", cache["cumulative_z_m"], z, tolerance)
        _cache_warning(issues, cells["cache_cumulative_north_m"], "累计北坐标", cache["cumulative_north_m"], north, tolerance)
        _cache_warning(issues, cells["cache_cumulative_east_m"], "累计东坐标", cache["cumulative_east_m"], east, tolerance)

    if any(item["severity"] == "error" for item in issues):
        raise GeometryInputError(issues)

    axis_value = user_settings.get("axis_azimuth_deg", meta.get("axis"))
    if _is_blank(axis_value):
        if math.hypot(east, north) <= 1e-12:
            raise GeometryInputError(issues + [_issue("error", "COINCIDENT_ENDPOINT", "首末点重合，必须显式设置剖面方位角", [])])
        axis = math.degrees(math.atan2(east, north)) % 360.0
        axis_method = "endpoint"
    else:
        axis = _number(axis_value, "settings.axis_azimuth_deg", issues)
        if axis is None or not 0 <= axis < 360:
            issues.append(_issue("error", "INVALID_AXIS_AZIMUTH", "剖面方位角必须在 [0, 360)", []))
            raise GeometryInputError(issues)
        axis_method = "explicit"
    axis_rad = math.radians(axis)

    nodes: List[Dict[str, Any]] = [{"index": 0, "east_m": 0.0, "north_m": 0.0, "z_m": 0.0,
                                    "x_m": 0.0, "offset_m": 0.0, "chainage_m": 0.0, "slant_chainage_m": 0.0}]
    ce = cn = cz = ch = cs = 0.0
    for record in records:
        comp = record["computed"]
        ce += comp["east_delta_m"]
        cn += comp["north_delta_m"]
        cz += comp["vertical_m"]
        ch += comp["horizontal_m"]
        cs += record["length_m"]
        nodes.append({"index": len(nodes), "east_m": ce, "north_m": cn, "z_m": cz,
                      "x_m": ce * math.sin(axis_rad) + cn * math.cos(axis_rad),
                      "offset_m": cn * math.sin(axis_rad) - ce * math.cos(axis_rad),
                      "chainage_m": ch, "slant_chainage_m": cs})

    intervals: List[Dict[str, Any]] = []
    for record in records:
        if not intervals or intervals[-1]["layer_id"] != record["layer_id"]:
            intervals.append({"id": f"I{len(intervals) + 1:04d}", "layer_id": record["layer_id"],
                              "start_node": record["start_node"], "end_node": record["end_node"],
                              "record_ids": [record["id"]], "lithology_name": record["lithology_name"],
                              "description": record["description"],
                              "source_cells": {"layer_id": [record["source_cells"]["layer_id"]],
                                               "lithology_name": [record["source_cells"].get("lithology_name", record["source_cells"]["description"])],
                                               "description": [record["source_cells"]["description"]]}})
        else:
            interval = intervals[-1]
            if (interval["lithology_name"] and record["lithology_name"]
                    and interval["lithology_name"] != record["lithology_name"]):
                issues.append(_issue("error", "LITHOLOGY_CONFLICT",
                                     f"连续层段 {record['layer_id']} 内存在不同明确岩性名称，请分层或修正层号",
                                     [interval["source_cells"]["lithology_name"][0],
                                      record["source_cells"]["lithology_name"]]))
            elif not interval["lithology_name"] and record["lithology_name"]:
                interval["lithology_name"] = record["lithology_name"]
            interval["end_node"] = record["end_node"]
            interval["record_ids"].append(record["id"])
            for key in ("layer_id", "description"):
                value = record["source_cells"][key]
                if value not in interval["source_cells"][key]:
                    interval["source_cells"][key].append(value)

    if any(item["severity"] == "error" for item in issues):
        raise GeometryInputError(issues)

    attitudes: List[Dict[str, Any]] = []
    last_attitude: Optional[Tuple[str, float, float]] = None
    for record in records:
        if record["dip_direction_deg"] is None:
            last_attitude = None
            continue
        key = (record["layer_id"], record["dip_direction_deg"], record["dip_angle_deg"])
        if key != last_attitude:
            attitudes.append({"id": f"A{len(attitudes) + 1:04d}", "record_id": record["id"],
                              "layer_id": record["layer_id"], "node": record["start_node"],
                              "dip_direction_deg": record["dip_direction_deg"], "dip_angle_deg": record["dip_angle_deg"],
                              "location_basis": "record_start_association"})
        last_attitude = key

    samples: List[Dict[str, Any]] = []
    for row, record in zip(input_rows, records):
        sample_id = _text(row["sample_id"])
        if not sample_id:
            if not _is_blank(row["sample_offset"]):
                issues.append(_issue("warning", "ORPHAN_SAMPLE_OFFSET", "提供了样品距离但没有样品编号，未创建样品", [record["source_cells"].get("sample_offset_m", "")]))
            continue
        sample_cells = {"sample_id": record["source_cells"]["sample_id"],
                        "sample_offset_m": record["source_cells"].get("sample_offset_m")}
        if _is_blank(row["sample_offset"]):
            samples.append({"id": sample_id, "record_id": record["id"], "layer_id": record["layer_id"],
                            "position": None, "location_status": "missing", "source_cells": sample_cells})
            continue
        offset_along = _number(row["sample_offset"], sample_cells["sample_offset_m"] or "", issues)
        if offset_along is None or not 0 <= offset_along <= record["length_m"]:
            issues.append(_issue("error", "INVALID_SAMPLE_OFFSET", "样品距测段起点必须在 [0, 斜距]", [sample_cells["sample_offset_m"] or ""]))
            continue
        ratio = offset_along / record["length_m"]
        start, end = nodes[record["start_node"]], nodes[record["end_node"]]
        position = {key: start[key] + (end[key] - start[key]) * ratio
                    for key in ("east_m", "north_m", "z_m", "x_m", "offset_m")}
        samples.append({"id": sample_id, "record_id": record["id"], "layer_id": record["layer_id"],
                        "position": position, "location_status": "explicit_offset", "source_cells": sample_cells})
    if any(item["severity"] == "error" for item in issues):
        raise GeometryInputError(issues)

    stations: List[Dict[str, Any]] = []
    current_end: Optional[str] = None
    previous_leg_id: Optional[str] = None
    for idx, (row, record) in enumerate(zip(input_rows, records)):
        if record["leg_id"] == previous_leg_id:
            continue
        parsed = _parse_leg(record["leg_id"])
        source_cell = record["source_cells"]["leg_id"]
        if parsed is None:
            issues.append(_issue("warning", "UNPARSEABLE_LEG_ID", f"导线段号“{record['leg_id']}”无法拆为起点-终点，使用自动站号", [source_cell]))
            start_id, end_id = f"AUTO-{record['start_node']}", f"AUTO-{record['end_node']}"
            current_end = None
        else:
            start_id, end_id = parsed
            if current_end is not None and start_id != current_end:
                issues.append(_issue("error", "DISCONNECTED_STATIONS", f"导线段 {record['leg_id']} 与前一可拆导线段不相接", [source_cell]))
            current_end = end_id
        if not stations or stations[-1]["node"] != record["start_node"]:
            stations.append({"id": start_id, "node": record["start_node"], "source_cell": source_cell})
        previous_leg_id = record["leg_id"]
    if records:
        last_cell = records[-1]["source_cells"]["leg_id"]
        final_id = current_end or f"AUTO-{len(records)}"
        stations.append({"id": final_id, "node": len(records), "source_cell": last_cell})
    if any(item["severity"] == "error" for item in issues):
        raise GeometryInputError(issues)

    endpoint = nodes[-1]
    return {
        "schema_version": SCHEMA_VERSION,
        "source": {"filename": source_path.name, "sha256": _sha256(source_path), "adapter": adapter, "sheet": meta["sheet"]},
        "project": {"name": meta["name"], "section_id": meta["section_id"]},
        "settings": {"axis_azimuth_deg": axis, "axis_method": axis_method,
                     "coordinate_system": "local_relative", "vertical_exaggeration": 1},
        "records": records, "nodes": nodes, "intervals": intervals, "stations": stations,
        "attitudes": attitudes, "samples": samples, "issues": issues,
        "summary": {"records": len(records), "intervals": len(intervals), "stations": len(stations),
                    "samples": len(samples), "located_samples": sum(s["position"] is not None for s in samples),
                    "total_slant_m": sum(r["length_m"] for r in records),
                    "total_horizontal_m": sum(r["computed"]["horizontal_m"] for r in records),
                    "total_vertical_m": endpoint["z_m"], "endpoint_east_m": endpoint["east_m"],
                    "endpoint_north_m": endpoint["north_m"], "axis_azimuth_deg": axis,
                    "projected_endpoint_m": endpoint["x_m"]},
    }


__all__ = ["GeometryInputError", "load_workbook_data"]
