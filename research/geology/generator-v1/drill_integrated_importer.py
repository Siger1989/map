"""Adapter for the 11-sheet integrated drilling workbook template.

Reads user input and cached formula values only; never evaluates workbook formulas.
"""
from __future__ import annotations

import hashlib
import json
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

from importer import GeometryInputError, _cell, _is_blank, _issue, _number, _text

INTEGRATED_SHEETS = {
    "钻孔基本信息", "回次表", "分层-原始记录", "分层-地质综合", "采样",
    "自定义测试项目", "样品测试结果", "孔深校正及弯曲度", "图签", "图面示意", "填写说明",
}

TABLE_HEADERS_V2 = {
    "回次表": ["序号","回次号","下界记录孔深_m","回次进尺_m","采取块数/原记","岩心长度_m","残留_m","处理后岩心长_m（手填）","回次采取率_pct（原填）","回次采取率_pct（计算）","孔深校正量_m","下界校正孔深_m","≤10cm岩心长度_m","≥10cm岩心长度_m","RQD_pct（原填）","备注","长度核对提示"],
    "分层-原始记录": ["序号","层号","起回次号","起位置_m","止回次号","止位置_m","换层累计孔深_m","分层进尺_m","分层岩心长_m","分层采取率_pct（原填）","分层采取率_pct（计算）","岩性名称","平均轴夹角_deg","真厚度_m（原填）","调整/说明","止回次岩心长_m","原表O列未命名值"],
    "分层-地质综合": ["序号","起回次号","起位置_m","止回次号","止位置_m","换层孔深_m","分层进尺_m","分层岩心长_m","分层采取率_pct（原填）","平均轴夹角_deg","真厚度_m（原填）"],
    "采样": ["序号","样品编号","起回次号","起位置_m","止回次号","止位置_m","孔深自_m","孔深至_m","进尺_m","岩心长度_m","采取率_pct（原填）","采取率_pct（计算）","重量数值（原单位未明）","重量单位（手填）","原表L列值（含义待确认）","原表M列未命名值/附加备注","原表N列未命名值/附加备注"],
    "自定义测试项目": ["项目代码","项目名称/元素","单位","检出限","分析方法","图中显示","显示顺序","备注"],
    "样品测试结果": ["样品编号","项目代码","结果原文","数值结果（可空）","结果标记","单位","检出限","分析方法","检测日期","报告编号","备注"],
}

TABLE_HEADERS_V3 = {
    "回次表": ["序号","回次号","下界记录孔深_m","回次进尺_m","采取块数/原记","岩心长度_m","残留_m","处理后岩心长_m（手填）","回次采取率_pct（计算）","孔深校正量_m","下界校正孔深_m","≤10cm岩心长度_m","≥10cm岩心长度_m","RQD_pct（原填）","备注","长度核对提示"],
    "分层-原始记录": ["序号","层号","起回次号","起位置_m","止回次号","止位置_m","换层累计孔深_m","分层进尺_m","分层岩心长_m","分层采取率_pct（计算）","岩性名称","平均轴夹角_deg","真厚度_m（原填）","调整/说明","止回次岩心长_m","原表O列未命名值"],
    "分层-地质综合": ["序号","起回次号","起位置_m","止回次号","止位置_m","换层孔深_m","分层进尺_m","分层岩心长_m","分层采取率_pct（计算）","平均轴夹角_deg","真厚度_m（原填）"],
    "采样": ["序号","样品编号","起回次号","起位置_m","止回次号","止位置_m","孔深自_m","孔深至_m","进尺_m","岩心长度_m","采取率_pct（计算）","重量数值（原单位未明）","重量单位（手填）","原表L列值（含义待确认）","原表M列未命名值/附加备注","原表N列未命名值/附加备注"],
    "自定义测试项目": TABLE_HEADERS_V2["自定义测试项目"],
    "样品测试结果": ["样品编号","项目代码","结果原文","数值结果（可空）","单位","检出限","分析方法","检测日期","报告编号","备注"],
}

TABLE_HEADERS_BY_VERSION = {"v2": TABLE_HEADERS_V2, "v3": TABLE_HEADERS_V3}

MAX_INPUT_ROWS = 10000


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _ref(sheet: str, address: str) -> str:
    return f"{sheet}!{address}"


def _source(ws, row: int, header_values: list[Any] | None = None) -> dict[str, Any]:
    raw = {}
    cells = {}
    if header_values:
        for col, label in enumerate(header_values, 1):
            if _is_blank(label):
                continue
            value = ws.cell(row, col).value
            raw[_text(label)] = _json_value(value)
            cells[_text(label)] = _ref(ws.title, ws.cell(row, col).coordinate)
    return {"sheet": ws.title, "row": row, "cells": cells, "raw": raw}


def _input_value(ws, row: int, col: int, issues: list[dict[str, Any]], *, field: str, required: bool = False):
    cell = ws.cell(row, col)
    value = cell.value
    ref = _ref(ws.title, cell.coordinate)
    if isinstance(value, str) and value.startswith("="):
        issues.append(_issue("error", "FORMULA_VALUE", f"{ref}（{field}）为公式；请填写原始值，适配器不会执行公式", [ref]))
        return None
    if required and _is_blank(value):
        issues.append(_issue("error", "MISSING_GEOMETRY", f"{ref} 缺少必需几何值", [ref]))
    return value


def _input_rows(ws, first: int, last: int, cols: list[int], formula_cols: set[int] | None = None):
    formula_cols = formula_cols or set()
    out = []
    for row in range(first, last + 1):
        populated = False
        for col in cols:
            if col in formula_cols:
                continue
            value = ws.cell(row, col).value
            if not _is_blank(value):
                populated = True
                break
        if populated:
            out.append(row)
    return out


def _number_at(ws, row: int, col: int, issues, field: str, required=False):
    value = _input_value(ws, row, col, issues, field=field, required=required)
    if isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("="):
        return None
    return _number(value, _ref(ws.title, ws.cell(row, col).coordinate), issues, required=False)


def _cached_number_at(ws, cached_ws, row: int, col: int, issues, field: str):
    cell=ws.cell(row,col)
    if isinstance(cell.value,str) and cell.value.startswith("="):
        cached_value=cached_ws.cell(row,col).value if cached_ws is not None else None
        return _number(cached_value,_ref(ws.title,cell.coordinate),issues,required=False) if not _is_blank(cached_value) else None
    return _number_at(ws,row,col,issues,field)


def _value(ws, addr: str, issues: list[dict[str, Any]], field: str, numeric=False):
    cell = ws[addr]
    value = _input_value(ws, cell.row, cell.column, issues, field=field)
    if numeric and not _is_blank(value):
        return _number(value, _ref(ws.title, addr), issues)
    return _json_value(value)


def _table_header(ws, row=5):
    return [_text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]


def _check_headers(ws, expected, row, issues, start_col=1):
    actual = [_text(ws.cell(row, start_col + i).value) for i in range(len(expected))]
    if actual != expected:
        issues.append(_issue("error", "UNSUPPORTED_TABLE_HEADERS",
            f"{ws.title} 表头与整合模板不一致；请恢复标准列名和顺序", [_ref(ws.title, ws.cell(row, start_col).coordinate)]))


def _detect_template_version(wb, issues):
    matches = {}
    for version, schema in TABLE_HEADERS_BY_VERSION.items():
        matches[version] = all(
            [_text(wb[sheet].cell(5, col).value) for col in range(1, len(headers)+1)] == headers
            for sheet, headers in schema.items() if sheet != "孔深校正及弯曲度"
        )
    if matches["v3"]:
        return "v3"
    if matches["v2"]:
        return "v2"
    issues.append(_issue("error","UNSUPPORTED_TEMPLATE_VERSION","整合模板列头不是完整的v2或v3格式，不能安全解释数据",[]))
    return "v2"


def _data_rows(ws, first, cols, formula_cols=None, issues=None):
    last = min(max(ws.max_row, first), MAX_INPUT_ROWS)
    if ws.max_row > MAX_INPUT_ROWS and issues is not None:
        issues.append(_issue("error", "INPUT_ROW_LIMIT", f"{ws.title} 数据超过安全上限 {MAX_INPUT_ROWS} 行", [_ref(ws.title, f"A{MAX_INPUT_ROWS + 1}")]))
    return _input_rows(ws, first, last, cols, formula_cols)


def _raw_table_rows(ws, first, last, header_row, formula_cols):
    headers = _table_header(ws, header_row)
    rows = []
    for row in _input_rows(ws, first, last, list(range(1, len(headers) + 1)), formula_cols):
        rows.append(_source(ws, row, headers))
    return rows


def _check_coverage(items, label, endpoint, issues):
    if not items:
        if label == "分层":
            issues.append(_issue("error", "EMPTY_LAYER_SHEET", "原始分层记录至少需要一条有效记录", ["分层-原始记录!6:25"]))
        return
    cursor = 0.0
    for item in items:
        top, bottom = item.get("top_m"), item.get("bottom_m")
        source = item.get("source", {}).get("cells", {})
        top_cell = source.get("顶深_m", source.get("derived_top_m", ""))
        bottom_cell = source.get("底深_m", "")
        if top is None or bottom is None:
            continue
        if bottom <= top:
            issues.append(_issue("error", "INVALID_INTERVAL", f"{label}底深必须大于顶深", [top_cell, bottom_cell]))
        if top < cursor - 1e-6:
            issues.append(_issue("error", "INTERVAL_OVERLAP", f"{label}区间重叠或倒序", [top_cell]))
        elif top > cursor + 1e-6:
            issues.append(_issue("error", "INTERVAL_GAP", f"{label}区间存在缺段：{cursor:g}–{top:g} m", [top_cell]))
        cursor = max(cursor, bottom)
        if top < -1e-6 or bottom > endpoint + 1e-6:
            issues.append(_issue("error", "INTERVAL_OUT_OF_RANGE", f"{label}区间超出0至终孔深度", [top_cell, bottom_cell]))
    if abs(cursor - endpoint) > 1e-6:
        issues.append(_issue("error", "INCOMPLETE_COVERAGE", f"{label}必须连续覆盖0至终孔深度；当前末深 {cursor:g} m，终孔 {endpoint:g} m", [items[-1].get("source", {}).get("cells", {}).get("底深_m", "")]))


def _basic_info(ws, cached_ws, issues):
    mapping = {
        "矿区": "B4", "线号": "D4", "钻孔编号": "F4", "项目名称": "H4",
        "开孔日期": "B5", "终孔日期": "D5", "终孔深度_m": "F5", "孔深基准": "H5",
        "设计方位角_deg": "B6", "倾角_deg": "D6", "X坐标_m": "F6", "Y坐标_m": "H6",
        "孔口标高_m": "B7", "坐标系": "D7", "角度定义": "F7", "比例尺分母": "H7", "备注": "B8",
    }
    numeric = {"终孔深度_m", "设计方位角_deg", "倾角_deg", "X坐标_m", "Y坐标_m", "孔口标高_m", "比例尺分母"}
    fields, cells = {}, {}
    for label, addr in mapping.items():
        fields[label] = _value(ws, addr, issues, label, numeric=label in numeric)
        cells[label] = _ref(ws.title, addr)
    summary = {}
    for row in range(12, 16):
        for label_col, value_col in ((1,2),(3,4),(5,6),(7,8)):
            label = _text(ws.cell(row, label_col).value)
            if not label:
                continue
            formula = ws.cell(row, value_col).value
            cached = cached_ws.cell(row, value_col).value if cached_ws else None
            summary[label] = {"value": _json_value(cached if isinstance(formula, str) and formula.startswith("=") else formula),
                              "formula": formula if isinstance(formula, str) and formula.startswith("=") else None,
                              "source": _ref(ws.title, ws.cell(row, value_col).coordinate)}
    return {"fields": fields, "source_cells": cells, "summary": summary}


def _metadata_sheet(ws, cached_ws, issues):
    labels = {
        "项目/单位": "B4", "图名": "B5", "拟编": "B6", "审核": "D6", "制图": "B7",
        "项目负责": "D7", "单位负责": "B8", "图号": "D8", "顺序号": "B9",
        "比例尺分母": "D9", "日期": "B10", "资料来源": "D10",
    }
    fields = {label: _value(ws, addr, issues, label) for label, addr in labels.items()}
    cells = {label: _ref(ws.title, addr) for label, addr in labels.items()}
    settings = []
    for row in range(15, 19):
        label = _text(ws.cell(row, 1).value)
        if not label:
            continue
        settings.append({"layout_item": label,
            "width_mm": _value(ws, f"B{row}", issues, f"{label}宽度_mm", numeric=True),
            "height_mm": _value(ws, f"C{row}", issues, f"{label}高度_mm", numeric=True),
            "order": _value(ws, f"D{row}", issues, f"{label}顺序", numeric=True),
            "description": _value(ws, f"E{row}", issues, f"{label}说明"),
            "source": {k: _ref(ws.title, f"{k}{row}") for k in "ABCDE"}})
    width_formula = ws["B19"].value
    width_cached = cached_ws["B19"].value if cached_ws is not None else None
    width_value = width_cached if isinstance(width_formula,str) and width_formula.startswith("=") else _value(ws,"B19",issues,"参考总宽_mm",numeric=True)
    summary = {"reference_total_width_mm": _json_value(width_value),
               "reference_total_width_formula": width_formula if isinstance(width_formula,str) and width_formula.startswith("=") else None,
               "row_height_mm": _value(ws, "D19", issues, "字段行高_mm", numeric=True),
               "source_cells": {"参考总宽_mm": _ref(ws.title,"B19"), "字段行高_mm": _ref(ws.title,"D19")}}
    return {"fields": fields, "source_cells": cells, "settings": settings, "settings_summary": summary}


def _depth_measurements(ws, cached_ws, issues):
    headers = _table_header(ws, 7)
    _check_headers(ws,["记录孔深_m","校测孔深_m","误差_m","误差率_pct","测量孔深_m","测量天顶角_deg","实测方位角_deg","测量方法","测量仪器"],7,issues,start_col=2)
    records = []
    seq = 0
    input_cols = list(range(2,11))
    footer_row = next((r for r in range(8,ws.max_row+1)
        if isinstance(ws.cell(r,1).value,str) and not ws.cell(r,1).value.startswith("=")
        and _is_blank(ws.cell(r,2).value)), None)
    data_end=min(footer_row-2 if footer_row is not None else ws.max_row,MAX_INPUT_ROWS)
    for row in _input_rows(ws, 8, max(7,data_end), input_cols, {1}):
        seq += 1
        raw = _source(ws, row, headers)
        seq_value = ws.cell(row,1).value
        if isinstance(seq_value, str) and seq_value.startswith("="):
            seq_value = seq
        elif _is_blank(seq_value):
            seq_value = seq
        fields = {}
        for idx, key in enumerate(("recorded_depth_m","checked_depth_m","error_m","error_percent","measurement_depth_m","zenith_deg","azimuth_deg","method","instrument"), start=2):
            value = _input_value(ws,row,idx,issues,field=key)
            fields[key] = _number(value,_ref(ws.title,ws.cell(row,idx).coordinate),issues) if key not in {"method","instrument"} and not _is_blank(value) and not (isinstance(ws.cell(row,idx).value,str) and ws.cell(row,idx).value.startswith("=")) else _text(value) if key in {"method","instrument"} else None
        fields.update({"sequence": _number(seq_value,_ref(ws.title,f"A{row}"),issues) if not _is_blank(seq_value) and not (isinstance(seq_value,str) and seq_value.startswith("=")) else seq,
                       "source": raw,
                       "derived": {"sequence": "row_order" if isinstance(ws.cell(row,1).value,str) and ws.cell(row,1).value.startswith("=") or _is_blank(ws.cell(row,1).value) else None}})
        records.append(fields)
    summary_label_row=(footer_row+1) if footer_row is not None else 28
    summary_value_row=footer_row+2
    signature_row=footer_row+3
    labels = [("应测次数",f"A{summary_label_row}",f"B{summary_value_row}"),("实测次数",f"C{summary_label_row}",f"D{summary_value_row}"),("超差次数",f"E{summary_label_row}",f"F{summary_value_row}"),
              ("弯曲度应测次数",f"G{summary_label_row}",f"H{summary_value_row}"),("弯曲度实测次数",f"I{summary_label_row}",f"J{summary_value_row}"),("弯曲度超差次数",f"K{summary_label_row}",f"L{summary_value_row}")]
    summary = {}
    for label, label_cell, value_cell in labels:
        raw_formula = ws[value_cell].value
        cached_value = cached_ws[value_cell].value if cached_ws is not None else None
        summary[label] = {"value": _json_value(cached_value if isinstance(raw_formula,str) and raw_formula.startswith("=") else raw_formula),
                          "formula": raw_formula if isinstance(raw_formula,str) and raw_formula.startswith("=") else None,
                          "source": _ref(ws.title,value_cell), "label_source": _ref(ws.title,label_cell)}
    signatures = {}
    for label, addr in (("记录人",f"B{signature_row}"),("记录日期",f"E{signature_row}"),("检查人",f"H{signature_row}"),("检查日期",f"K{signature_row}")):
        signatures[label] = {"value": _json_value(ws[addr].value), "source": _ref(ws.title,addr)}
    return {"records": records, "summary": summary, "signatures": signatures}


def load_integrated_drill_workbook(path: str | Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("读取整合钻孔 XLSX 需要 openpyxl") from exc
    path = Path(path)
    wb = load_workbook(path, read_only=False, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    issues: list[dict[str, Any]] = []
    missing_sheets = sorted(INTEGRATED_SHEETS - set(wb.sheetnames))
    if missing_sheets:
        wb.close(); cached.close()
        raise GeometryInputError([_issue("error","MISSING_INTEGRATED_SHEET",f"整合模板缺少工作表：{', '.join(missing_sheets)}",[])])
    template_version=_detect_template_version(wb,issues)
    table_headers=TABLE_HEADERS_BY_VERSION[template_version]
    basic_ws = wb["钻孔基本信息"]
    basic = _basic_info(basic_ws,cached["钻孔基本信息"],issues)
    basic_fields = basic["fields"]
    hole_id = _text(basic_fields.get("钻孔编号"))
    endpoint_raw = basic_fields.get("终孔深度_m")
    endpoint = endpoint_raw if isinstance(endpoint_raw,(int,float)) and not isinstance(endpoint_raw,bool) else None
    basis = _text(basic_fields.get("孔深基准"))
    if not hole_id:
        issues.append(_issue("error","MISSING_HOLE_ID","钻孔编号必填",[basic["source_cells"]["钻孔编号"]]))
    if endpoint is None:
        issues.append(_issue("error","MISSING_ENDPOINT","终孔深度_m 必填且必须为数值",[basic["source_cells"]["终孔深度_m"]]))
    elif endpoint <= 0:
        issues.append(_issue("error","INVALID_ENDPOINT","终孔深度必须大于0",[basic["source_cells"]["终孔深度_m"]]))
    if basis not in ("原始沿孔深","校正沿孔深"):
        issues.append(_issue("error","INVALID_DEPTH_BASIS","孔深基准必须明确为原始沿孔深或校正沿孔深",[basic["source_cells"]["孔深基准"]]))
    endpoint_for_checks = endpoint if endpoint is not None and endpoint > 0 else 0.0

    # Element directory and long-form results. Codes remain user-defined.
    catalog_ws = wb["自定义测试项目"]
    _check_headers(catalog_ws,table_headers[catalog_ws.title],5,issues)
    catalog_headers = _table_header(catalog_ws,5)
    analysis_items, units, catalog = [], {}, {}
    for row in _data_rows(catalog_ws,6,list(range(1,9)),issues=issues):
        vals = [_input_value(catalog_ws,row,c,issues,field=catalog_headers[c-1]) for c in range(1,9)]
        code = _text(vals[0]); name = _text(vals[1]); unit = _text(vals[2])
        source = _source(catalog_ws,row,catalog_headers)
        if not code:
            issues.append(_issue("error","MISSING_ANALYSIS_CODE","自定义测试项目代码必填",[_ref(catalog_ws.title,f"A{row}")]))
            continue
        if code in catalog:
            issues.append(_issue("error","DUPLICATE_ANALYSIS_CODE",f"测试项目代码重复：{code}",[_ref(catalog_ws.title,f"A{row}"),catalog[code]["source"]["cells"].get("项目代码","")]))
            continue
        if not name:
            issues.append(_issue("error","MISSING_ANALYSIS_NAME",f"测试项目 {code} 缺少名称",[_ref(catalog_ws.title,f"B{row}")]))
        detect = _input_value(catalog_ws,row,4,issues,field="检出限")
        method = _text(_input_value(catalog_ws,row,5,issues,field="分析方法"))
        show = _text(_input_value(catalog_ws,row,6,issues,field="图中显示"))
        if show and show not in ("是", "否"):
            issues.append(_issue("error","INVALID_ANALYSIS_SHOW",f"项目 {code} 的图中显示只允许“是”或“否”",[_ref(catalog_ws.title,f"F{row}")]))
        order = _number_at(catalog_ws,row,7,issues,"显示顺序")
        item = {"code":code,"name":name,"unit":unit,"order":order,"detection_limit":_json_value(detect),"method":method,
                "show": show != "否","notes":_text(vals[7]),"source":source}
        catalog[code]=item; units[code]=unit; analysis_items.append(item)
    analysis_items.sort(key=lambda x:(x["order"] is None,x["order"] if x["order"] is not None else math.inf,x["code"]))

    # Layers use raw rows only. Top depth is explicitly derived from the prior bottom.
    synth_ws=wb["分层-地质综合"]
    _check_headers(synth_ws,table_headers[synth_ws.title],5,issues)
    layer_ws = wb["分层-原始记录"]; layer_headers = _table_header(layer_ws,5)
    _check_headers(layer_ws,table_headers[layer_ws.title],5,issues)
    layers=[]; previous_bottom=0.0; previous_bottom_source=None; layer_ids=set()
    layer_core_col=9; layer_length_col=8
    layer_original_recovery_col=10 if template_version=="v2" else None
    layer_computed_recovery_col=11 if template_version=="v2" else 10
    layer_last_col=17 if template_version=="v2" else 16
    formula_cols={1,layer_computed_recovery_col}
    for row in _data_rows(layer_ws,6,list(range(2,layer_last_col+1)),formula_cols,issues):
        source=_source(layer_ws,row,layer_headers)
        ident=_text(_input_value(layer_ws,row,2,issues,field="层号",required=True))
        bottom=_number_at(layer_ws,row,7,issues,"换层累计孔深_m",required=True)
        length=_number_at(layer_ws,row,layer_length_col,issues,"分层进尺_m")
        core=_number_at(layer_ws,row,layer_core_col,issues,"分层岩心长_m")
        raw_recovery=_number_at(layer_ws,row,layer_original_recovery_col,issues,"分层采取率_pct（原填）") if layer_original_recovery_col else None
        computed_recovery=(core/length*100 if core is not None and length is not None and length>0 else None)
        computed_recovery_source=_cached_number_at(layer_ws,cached["分层-原始记录"],row,layer_computed_recovery_col,issues,"分层采取率_pct（计算）")
        if not ident: issues.append(_issue("error","MISSING_ID","层号不能为空",[_ref(layer_ws.title,f"B{row}")]))
        elif ident in layer_ids: issues.append(_issue("error","DUPLICATE_ID",f"层号重复：{ident}",[_ref(layer_ws.title,f"B{row}")]))
        layer_ids.add(ident)
        top=previous_bottom if bottom is not None else None
        top_source=previous_bottom_source
        if bottom is not None and length is not None and top is not None and abs((bottom-top)-length)>0.01:
            issues.append(_issue("warning","LAYER_ADVANCE_MISMATCH","分层进尺与连续底深差值不一致；沿用底深推导几何",[_ref(layer_ws.title,f"H{row}"),_ref(layer_ws.title,f"G{row}")]))
        if core is not None and core < 0: issues.append(_issue("error","NEGATIVE_CORE_LENGTH","分层岩心长不能为负数",[_ref(layer_ws.title,f"I{row}")]))
        if core is not None and length is not None and length>0 and core>length+1e-6: issues.append(_issue("warning","CORE_EXCEEDS_INTERVAL","分层岩心长超过填入分层进尺，请核实",[_ref(layer_ws.title,f"I{row}"),_ref(layer_ws.title,f"H{row}")]))
        lith_col=12 if template_version=="v2" else 11
        adjustment_col=15 if template_version=="v2" else 14
        stop_core_col=16 if template_version=="v2" else 15
        unknown_col=17 if template_version=="v2" else 16
        lith=_text(_input_value(layer_ws,row,lith_col,issues,field="岩性名称"))
        # The last source columns are retained, but no geology/true-thickness inference is made.
        item={"id":ident,"top_m":top,"bottom_m":bottom,"thickness_m":bottom-top if bottom is not None and top is not None else None,
              "core_m":core,"recovery_percent":computed_recovery,"recovery_original_percent":raw_recovery,"recovery_computed_percent_source":computed_recovery_source,"description":_text(_input_value(layer_ws,row,adjustment_col,issues,field="调整/说明")),
              "lithology_name":lith,"material_code":"","pattern_id":None,
              "source":source | {"cells":source["cells"] | {"顶深_m":top_source or "derived:first-row-zero","底深_m":_ref(layer_ws.title,f"G{row}")},
                                  "derived":{"top_m":{"rule":"first row 0; subsequent row previous explicit cumulative bottom","source":top_source,"value":top}}}}
        item["true_thickness_m_raw"]=_json_value(layer_ws.cell(row,14 if template_version=="v2" else 13).value)
        item["mean_axis_angle_deg"]=_json_value(layer_ws.cell(row,13 if template_version=="v2" else 12).value)
        item["raw_layer_notes"]={"adjustment":_json_value(layer_ws.cell(row,adjustment_col).value),"stop_run_core_m":_json_value(layer_ws.cell(row,stop_core_col).value),"source_O_unlabeled":_json_value(layer_ws.cell(row,unknown_col).value)}
        layers.append(item)
        if bottom is not None: previous_bottom=bottom; previous_bottom_source=_ref(layer_ws.title,f"G{row}")
    if not layers:
        issues.append(_issue("error","EMPTY_LAYER_SHEET","原始分层记录至少需要一条有效记录",["分层-原始记录!6:25"]))
    if endpoint is not None:
        _check_coverage(layers,"分层",endpoint_for_checks,issues)

    # Rounds: use the manually entered processed core length, and keep other source measurements raw.
    round_ws=wb["回次表"]; round_headers=_table_header(round_ws,5)
    _check_headers(round_ws,table_headers[round_ws.title],5,issues)
    turns=[]; previous_bottom=0.0; previous_bottom_source=None; round_ids=set()
    round_original_recovery_col=9 if template_version=="v2" else None
    round_computed_recovery_col=10 if template_version=="v2" else 9
    round_last_col=17 if template_version=="v2" else 16
    for row in _data_rows(round_ws,6,list(range(2,round_last_col+1)),{1,round_computed_recovery_col,round_last_col},issues):
        source=_source(round_ws,row,round_headers)
        ident=_text(_input_value(round_ws,row,2,issues,field="回次号",required=True))
        bottom=_number_at(round_ws,row,3,issues,"下界记录孔深_m",required=True)
        advance=_number_at(round_ws,row,4,issues,"回次进尺_m")
        core=_number_at(round_ws,row,8,issues,"处理后岩心长_m（手填）")
        raw_recovery=_number_at(round_ws,row,round_original_recovery_col,issues,"回次采取率_pct（原填）") if round_original_recovery_col else None
        computed_recovery_source=_cached_number_at(round_ws,cached["回次表"],row,round_computed_recovery_col,issues,"回次采取率_pct（计算）")
        if not ident: issues.append(_issue("error","MISSING_ID","回次号不能为空",[_ref(round_ws.title,f"B{row}")]))
        elif ident in round_ids: issues.append(_issue("error","DUPLICATE_ID",f"回次号重复：{ident}",[_ref(round_ws.title,f"B{row}")]))
        round_ids.add(ident)
        top=previous_bottom if bottom is not None else None
        top_source=previous_bottom_source
        derived_advance=bottom-top if bottom is not None and top is not None else None
        if advance is not None and derived_advance is not None and abs(advance-derived_advance)>0.01:
            issues.append(_issue("warning","ROUND_ADVANCE_MISMATCH","回次进尺与连续底深差值不一致；沿用底深推导几何",[_ref(round_ws.title,f"D{row}"),_ref(round_ws.title,f"C{row}")]))
        if core is not None and core < 0: issues.append(_issue("error","NEGATIVE_CORE_LENGTH","处理后岩心长不能为负数",[_ref(round_ws.title,f"H{row}")]))
        computed_recovery=(core/advance*100 if core is not None and advance is not None and advance>0 else None)
        item={"id":ident,"top_m":top,"bottom_m":bottom,"advance_m":derived_advance,"core_m":core,"recovery_percent":computed_recovery,"recovery_original_percent":raw_recovery,"recovery_computed_percent_source":computed_recovery_source,
              "source":source | {"cells":source["cells"] | {"顶深_m":top_source or "derived:first-row-zero","底深_m":_ref(round_ws.title,f"C{row}"),"岩心长_m":_ref(round_ws.title,f"H{row}")},
                                  "derived":{"top_m":{"rule":"first row 0; subsequent row previous explicit bottom","source":top_source,"value":top},"advance_m":{"rule":"bottom_m - top_m","value":derived_advance}}}}
        item["source_metrics"]={"raw_core_length_m":_json_value(round_ws.cell(row,6).value),"raw_recovery_percent":_json_value(round_ws.cell(row,round_original_recovery_col).value) if round_original_recovery_col else None,"computed_recovery_formula":_json_value(round_ws.cell(row,round_computed_recovery_col).value),"depth_correction_m":_json_value(round_ws.cell(row,11 if template_version=="v2" else 10).value),"corrected_bottom_depth_m":_json_value(round_ws.cell(row,12 if template_version=="v2" else 11).value),"rqd_raw":_json_value(round_ws.cell(row,15 if template_version=="v2" else 14).value)}
        turns.append(item)
        if bottom is not None: previous_bottom=bottom; previous_bottom_source=_ref(round_ws.title,f"C{row}")
    if turns and endpoint is not None: _check_coverage(turns,"回次",endpoint_for_checks,issues)

    # Samples plus configurable analytical results.
    sample_ws=wb["采样"]; sample_headers=_table_header(sample_ws,5)
    _check_headers(sample_ws,table_headers[sample_ws.title],5,issues)
    samples=[]; sample_by_id={}
    sample_original_recovery_col=11 if template_version=="v2" else None
    sample_computed_recovery_col=12 if template_version=="v2" else 11
    sample_last_col=17 if template_version=="v2" else 16
    for row in _data_rows(sample_ws,6,list(range(2,sample_last_col+1)),{1,sample_computed_recovery_col},issues):
        source=_source(sample_ws,row,sample_headers)
        ident=_text(_input_value(sample_ws,row,2,issues,field="样品编号",required=True))
        top=_number_at(sample_ws,row,7,issues,"孔深自_m",required=True)
        bottom=_number_at(sample_ws,row,8,issues,"孔深至_m",required=True)
        entered_length=_number_at(sample_ws,row,9,issues,"进尺_m")
        core=_number_at(sample_ws,row,10,issues,"岩心长度_m")
        raw_recovery=_number_at(sample_ws,row,sample_original_recovery_col,issues,"采取率_pct（原填）") if sample_original_recovery_col else None
        computed_recovery_source=_cached_number_at(sample_ws,cached["采样"],row,sample_computed_recovery_col,issues,"采取率_pct（计算）")
        if not ident: issues.append(_issue("error","MISSING_ID","样品编号不能为空",[_ref(sample_ws.title,f"B{row}")]))
        if ident in sample_by_id:
            issues.append(_issue("error","DUPLICATE_ID",f"样品编号重复：{ident}",[sample_by_id[ident]["source"]["cells"].get("样品编号",""),_ref(sample_ws.title,f"B{row}")]))
        if bottom is not None and top is not None and bottom<=top:
            issues.append(_issue("error","INVALID_INTERVAL","样品孔深至必须大于孔深自",[_ref(sample_ws.title,f"G{row}"),_ref(sample_ws.title,f"H{row}")]))
        if top is not None and (top<0 or endpoint is not None and bottom is not None and bottom>endpoint+1e-6):
            issues.append(_issue("error","SAMPLE_OUT_OF_RANGE","样品区间超出孔深范围",[_ref(sample_ws.title,f"G{row}"),_ref(sample_ws.title,f"H{row}")]))
        length=bottom-top if bottom is not None and top is not None else None
        if entered_length is not None and length is not None and abs(entered_length-length)>0.01:
            issues.append(_issue("warning","SAMPLE_ADVANCE_MISMATCH","样品进尺与深度区间差值不一致；沿用深度区间",[_ref(sample_ws.title,f"I{row}"),_ref(sample_ws.title,f"G{row}"),_ref(sample_ws.title,f"H{row}")]))
        assay_init={code:None for code in catalog}; raw_init={code:"" for code in catalog}
        computed_recovery=(core/entered_length*100 if core is not None and entered_length is not None and entered_length>0 else None)
        item={"id":ident,"top_m":top,"bottom_m":bottom,"length_m":length,"core_m":core,"recovery_percent":computed_recovery,"recovery_original_percent":raw_recovery,"recovery_computed_percent_source":computed_recovery_source,
              "assays":assay_init,"assay_raw":raw_init,"assay_records":[],"position":None,
              "source":source | {"cells":source["cells"] | {"顶深_m":_ref(sample_ws.title,f"G{row}"),"底深_m":_ref(sample_ws.title,f"H{row}")}}}
        item["sample_details"]={"start_round":_json_value(sample_ws.cell(row,3).value),"start_position_m":_json_value(sample_ws.cell(row,4).value),"end_round":_json_value(sample_ws.cell(row,5).value),"end_position_m":_json_value(sample_ws.cell(row,6).value),"raw_recovery_percent":raw_recovery,"computed_recovery_percent_source":computed_recovery_source,"weight_value":_json_value(sample_ws.cell(row,13 if template_version=="v2" else 12).value),"weight_unit":_json_value(sample_ws.cell(row,14 if template_version=="v2" else 13).value),"source_L_unconfirmed":_json_value(sample_ws.cell(row,15 if template_version=="v2" else 14).value),"unlabeled_M":_json_value(sample_ws.cell(row,16 if template_version=="v2" else 15).value),"unlabeled_N":_json_value(sample_ws.cell(row,17 if template_version=="v2" else 16).value)}
        samples.append(item)
        if ident and ident not in sample_by_id: sample_by_id[ident]=item
    if samples and endpoint is not None:
        ordered=sorted(samples,key=lambda s:(s["top_m"] if s["top_m"] is not None else math.inf,s["bottom_m"] if s["bottom_m"] is not None else math.inf))
        for a,b in zip(ordered,ordered[1:]):
            if a["bottom_m"] is not None and b["top_m"] is not None and b["top_m"]<a["bottom_m"]-1e-9:
                issues.append(_issue("warning","OVERLAPPING_SAMPLES",f"样品区间存在重叠，将分轨显示：{a['id']} / {b['id']}",[a["source"]["cells"].get("顶深_m",""),b["source"]["cells"].get("顶深_m","")]))
                break

    result_ws=wb["样品测试结果"]; result_headers=_table_header(result_ws,5); seen_assays=set()
    _check_headers(result_ws,table_headers[result_ws.title],5,issues)
    result_last_col=11 if template_version=="v2" else 10
    for row in _data_rows(result_ws,6,list(range(1,result_last_col+1)),issues=issues):
        source=_source(result_ws,row,result_headers)
        sample_id=_text(_input_value(result_ws,row,1,issues,field="样品编号",required=True))
        code=_text(_input_value(result_ws,row,2,issues,field="项目代码",required=True))
        raw_value=_input_value(result_ws,row,3,issues,field="结果原文")
        raw_text=raw_value if isinstance(raw_value,str) else "" if _is_blank(raw_value) else _text(raw_value)
        numeric=_number_at(result_ws,row,4,issues,"数值结果")
        unit_col=6 if template_version=="v2" else 5
        detection_col=7 if template_version=="v2" else 6
        method_col=8 if template_version=="v2" else 7
        date_col=9 if template_version=="v2" else 8
        report_col=10 if template_version=="v2" else 9
        notes_col=11 if template_version=="v2" else 10
        unit=_text(_input_value(result_ws,row,unit_col,issues,field="单位"))
        if sample_id not in sample_by_id:
            issues.append(_issue("error","UNKNOWN_SAMPLE",f"测试结果引用未定义样品：{sample_id}",[_ref(result_ws.title,f"A{row}")]))
        if code not in catalog:
            issues.append(_issue("error","UNKNOWN_ANALYSIS_CODE",f"测试结果引用未定义项目代码：{code}",[_ref(result_ws.title,f"B{row}")]))
        if not code:
            issues.append(_issue("error","MISSING_ANALYSIS_CODE","测试结果项目代码必填",[_ref(result_ws.title,f"B{row}")]))
        key=(sample_id,code)
        if key in seen_assays:
            issues.append(_issue("error","DUPLICATE_SAMPLE_ANALYSIS",f"样品 {sample_id} 的项目 {code} 重复",[_ref(result_ws.title,f"A{row}"),_ref(result_ws.title,f"B{row}")]))
        seen_assays.add(key)
        if code in catalog and unit and catalog[code]["unit"] and unit!=catalog[code]["unit"]:
            issues.append(_issue("error","ANALYSIS_UNIT_CONFLICT",f"项目 {code} 的结果单位“{unit}”与目录单位“{catalog[code]['unit']}”不一致",[_ref(result_ws.title,result_ws.cell(row,unit_col).coordinate),catalog[code]["source"]["cells"].get("单位","")]))
        detail={"code":code,"result_raw":raw_text,"value":numeric,"flag":_text(result_ws.cell(row,5).value) if template_version=="v2" else None,"unit":unit,
                "detection_limit":_json_value(result_ws.cell(row,detection_col).value),"method":_text(result_ws.cell(row,method_col).value),"test_date":_json_value(result_ws.cell(row,date_col).value),
                "report_number":_text(result_ws.cell(row,report_col).value),"notes":_text(result_ws.cell(row,notes_col).value),"source":source}
        if sample_id in sample_by_id:
            sample=sample_by_id[sample_id]
            if code in catalog:
                sample["assays"][code]=numeric
                sample["assay_raw"][code]=raw_text
            sample["assay_records"].append(detail)

    # Preserve separate manual geology synthesis rows as source records only.
    synth_rows=_raw_table_rows(synth_ws,6,min(max(synth_ws.max_row,6),MAX_INPUT_ROWS),5,{1})
    for record in synth_rows:
        formulas={k:v for k,v in record["raw"].items() if isinstance(v,str) and v.startswith("=")}
        input_values=[v for k,v in record["raw"].items() if k!="序号" and not (isinstance(v,str) and v.startswith("=")) and not _is_blank(v)]
        if input_values:
            issues.append(_issue("warning","SYNTHESIS_NOT_DRAWN","分层-地质综合是独立人工汇总页；当前绘图几何仅取分层-原始记录，不会覆盖或融合本页",[f"分层-地质综合!B{record['row']}:K{record['row']}"]))
    depth=_depth_measurements(wb["孔深校正及弯曲度"],cached["孔深校正及弯曲度"],issues)
    title_block=_metadata_sheet(wb["图签"],cached["图签"],issues)

    # Material pattern mapping remains exact-name based, as in the legacy adapter.
    patterns=json.loads((Path(__file__).resolve().parent/"templates"/"drill-patterns.json").read_text(encoding="utf-8"))
    codes, materials=patterns["codes"],patterns["materials"]
    for layer in layers:
        lith=layer.get("lithology_name",""); expected=materials.get(lith)
        if expected:
            layer["material_code"]=expected; layer["pattern_id"]=codes[expected]
        else:
            issues.append(_issue("warning","PENDING_PATTERN",f"岩性“{lith or '未提供'}”无精确花纹映射，保留留白待配置",[layer.get("source",{}).get("cells",{}).get("岩性名称","")]))
    source={"filename":path.name,"sha256":_sha(path),"adapter":f"drill-integrated-template-{template_version}","template_version":template_version,"sheet_names":list(wb.sheetnames)}
    wb.close(); cached.close()
    if any(i.get("severity")=="error" for i in issues):
        raise GeometryInputError(issues)
    project={"name":_text(basic_fields.get("项目名称")),"hole_id":hole_id,"analysis_units":units,"analysis_items":analysis_items,
             "source_cells":basic["source_cells"]}
    meta={"hole_id":hole_id,"endpoint_m":endpoint,"depth_basis":basis}
    summary={"turns":len(turns),"layers":len(layers),"samples":len(samples),"endpoint_m":endpoint,"pending_patterns":sum(1 for x in layers if not x.get("material_code"))}
    return {"drawing_type":"drill","schema_version":"drill-integrated-1.0","source":source,"project":project,"meta":meta,
            "basic_info":basic,"turns":turns,"layers":layers,"samples":samples,"structures":[],"issues":issues,"summary":summary,
            "depth_measurements":depth,"title_block":title_block,"source_tables":{"分层-地质综合":synth_rows}}
