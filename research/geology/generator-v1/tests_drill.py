"""Focused behavioral tests for drill workbook validation and geometry."""
import sys, tempfile, unittest
from pathlib import Path

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))
from drill_importer import load_drill_workbook
from drill_renderer import _wrap_text, depth_y, render_drill_svg, sample_lanes
from drill_pipeline import _appendix, _raster_view, _hd_raster_view, _hd_dimension_check
from importer import GeometryInputError

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook=None

PARAMS=[('模板类型','钻孔柱状图'),('模板版本','1.0'),('项目名称','测试项目'),('钻孔编号','BH-X17'),('长度单位','m'),('孔深基准','原始沿孔深'),('终孔深度_m',13.5),('Au单位','ppm'),('Pb单位','ppm'),('Zn单位','ppm')]
LHEAD=['层号','顶深_m','底深_m','岩性名称','岩性描述','岩心长_m','花纹代码']
THEAD=['回次号','顶深_m','底深_m','岩心长_m']
SHEAD=['样品编号','顶深_m','底深_m','岩心长_m','Au','Pb','Zn']
PHEAD=['孔深_m','孔径_mm']

def make_book(path, *, layers=None, turns=None, samples=None, structures=None):
    wb=Workbook(); p=wb.active; p.title='项目'; p.append(['参数','值'])
    for x in PARAMS:p.append(x)
    for name,head,records in [('分层',LHEAD,layers if layers is not None else [['L-a',0,6,'砂岩夹粉砂岩','上段描述',5,''],['L-b',6,13.5,'未知岩性','保留原文？/说明',None,'']]),('回次',THEAD,turns if turns is not None else [['R-9',0,4,3.6],['R-10',4,13.5,8.0]]),('样品',SHEAD,samples if samples is not None else [['S-21',2,8,5,None,None,None],['S-22',5,9,None,1.25,None,None],['S-23',9,11,None,None,None,None]]),('孔径',PHEAD,structures or [])]:
        ws=wb.create_sheet(name); ws.append(head)
        for row in records:ws.append(row)
    wb.save(path); wb.close()

@unittest.skipIf(Workbook is None,'openpyxl is required')
class DrillWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)/'input.xlsx'; make_book(self.path)
    def tearDown(self):self.tmp.cleanup()
    def codes(self,fn):
        with self.assertRaises(GeometryInputError) as ctx:fn()
        return {i['code'] for i in ctx.exception.issues}
    def test_sample_recovery_percent_stays_with_its_number_when_wrapped(self):
        lines=_wrap_text('样品回收率 12.34%',30,10)
        self.assertIn('12.34%',lines)
    def test_dynamic_records_source_cells_null_assays_and_geometry(self):
        d=load_drill_workbook(self.path)
        self.assertEqual(d['meta']['hole_id'],'BH-X17'); self.assertEqual(d['meta']['endpoint_m'],13.5)
        self.assertEqual((len(d['layers']),len(d['turns']),len(d['samples'])),(2,2,3))
        self.assertEqual(d['layers'][0]['bottom_m']-d['layers'][0]['top_m'],6)
        self.assertEqual(d['layers'][0]['source']['cells']['顶深_m'],'分层!B2')
        self.assertIsNone(d['samples'][0]['assays']['Au'])
        self.assertEqual(d['summary']['pending_patterns'],1)
        self.assertIn('PENDING_PATTERN',{x['code'] for x in d['issues']})
        svg,audit=render_drill_svg(d)
        self.assertAlmostEqual(depth_y(13.5,scale=audit['scale_units_per_meter'])-depth_y(0,scale=audit['scale_units_per_meter']),13.5*audit['scale_units_per_meter'])
        self.assertIn('BH-X17',svg); self.assertIn('未提供',svg); self.assertNotIn('<image',svg)
        self.assertEqual(audit['counts'],{'layers':2,'turns':2,'samples':3})
        appendix=_appendix(d)
        self.assertIn('<th>Au (ppm)</th>',appendix); self.assertIn('来源单元格',appendix)
        self.assertIn('1.25',appendix)
    def test_overlapping_samples_are_legal_and_separate_tracks(self):
        d=load_drill_workbook(self.path); lanes,n=sample_lanes(d['samples'])
        self.assertEqual(n,2); self.assertNotEqual(lanes['S-21'],lanes['S-22'])
        self.assertIn('OVERLAPPING_SAMPLES',{x['code'] for x in d['issues']})
    def test_rejects_bad_coverage_and_out_of_order_rows(self):
        make_book(self.path,layers=[['L1',0,8,'x','','', ''],['L2',7,13.5,'y','','','']])
        self.assertTrue({'INTERVAL_OVERLAP'} & self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path,layers=[['L1',0,5,'x','','',''],['L2',7,13.5,'y','','','']])
        self.assertTrue({'INTERVAL_GAP'} & self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path,layers=[['L1',0,5,'x','','',''],['L2',5,13.5,'y','','','']],turns=[['R2',5,13.5,8],['R1',0,5,4]])
        self.assertTrue({'INTERVAL_GAP','INTERVAL_OVERLAP'} & self.codes(lambda:load_drill_workbook(self.path)))
    def test_rejects_missing_layers_negative_core_and_nonpositive_diameter(self):
        make_book(self.path,layers=[])
        self.assertTrue({'EMPTY_LAYER_SHEET'} & self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path,layers=[['L1',0,13.5,'x','',-1,'']])
        self.assertTrue({'NEGATIVE_CORE_LENGTH'} & self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path,structures=[[0,0]])
        self.assertTrue({'INVALID_DIAMETER'} & self.codes(lambda:load_drill_workbook(self.path)))
    def test_rejects_duplicate_project_formula_and_formula_data_cells(self):
        wb=load_workbook(self.path); ws=wb['项目']; ws.append(['钻孔编号','BH-DUP']); ws['B7']='=1+1'; wb.save(self.path); wb.close()
        codes=self.codes(lambda:load_drill_workbook(self.path))
        self.assertIn('DUPLICATE_PROJECT_PARAMETER',codes); self.assertIn('FORMULA_VALUE',codes)
        make_book(self.path); wb=load_workbook(self.path); wb['样品']['E2']='=1+1'; wb.save(self.path); wb.close()
        self.assertIn('FORMULA_VALUE',self.codes(lambda:load_drill_workbook(self.path)))
    def test_rejects_bad_sample_depth_and_code_mismatch(self):
        make_book(self.path,samples=[['S-x',13,14,None,None,None,None]])
        self.assertIn('SAMPLE_OUT_OF_RANGE',self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path,layers=[['L1',0,13.5,'砂岩夹粉砂岩','',None,'D001']])
        self.assertIn('PATTERN_MATERIAL_CONFLICT',self.codes(lambda:load_drill_workbook(self.path)))
    def test_duplicate_ids_and_missing_required_column_are_reported(self):
        make_book(self.path,samples=[['S-x',1,2,None,None,None,None],['S-x',3,4,None,None,None,None]])
        self.assertIn('DUPLICATE_ID',self.codes(lambda:load_drill_workbook(self.path)))
        make_book(self.path); wb=load_workbook(self.path); ws=wb['分层']; ws.delete_cols(2); wb.save(self.path); wb.close()
        codes=self.codes(lambda:load_drill_workbook(self.path))
        self.assertIn('MISSING_COLUMN',codes)
    def test_raster_view_keeps_equal_square_dimensions_and_viewbox(self):
        source=Path(self.tmp.name)/'square.svg'; dest=Path(self.tmp.name)/'raster.svg'
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="20000" height="20000" viewBox="0 0 20000 20000"></svg>',encoding='utf-8')
        meta=_raster_view(source,dest)
        text=dest.read_text(encoding='utf-8')
        self.assertEqual(meta['raster_view_size'],[8000,8000])
        self.assertIn('width="8000" height="8000" viewBox="0 0 20000 20000"',text)
    def test_complete_hd_uses_fractional_svg_view_and_keeps_full_native_viewbox(self):
        source=Path(self.tmp.name)/'elongated.svg'; dest=Path(self.tmp.name)/'hd-raster-view.svg'
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="2484" height="6267" viewBox="0 0 2484 6267"><rect width="2484" height="6267"/></svg>',encoding='utf-8')
        meta=_hd_raster_view(source,dest)
        self.assertGreater(meta['planned_scale_factor'],1)
        self.assertLess(meta['planned_scale_factor'],4)
        self.assertEqual(meta['hd_raster_view_size'][1],12000)
        self.assertLessEqual(meta['hd_raster_view_size'][0]*meta['hd_raster_view_size'][1],64_000_000)
        view=dest.read_text(encoding='utf-8')
        self.assertIn('width="%d" height="12000" viewBox="0 0 2484 6267"'%meta['hd_raster_view_size'][0],view)
        raster={'png_width_px':meta['hd_raster_view_size'][0],'png_height_px':meta['hd_raster_view_size'][1]}
        result=_hd_dimension_check(meta,raster,{'png_width_px':2484,'png_height_px':6267})
        self.assertTrue(result['passed']); self.assertTrue(result['full_viewBox_preserved']); self.assertTrue(result['png_exceeds_preview_resolution'])
        unchanged=_hd_dimension_check(meta,raster,raster)
        self.assertFalse(unchanged['png_exceeds_preview_resolution'])
        with self.assertRaises(RuntimeError):
            _hd_dimension_check(meta,{'png_width_px':2484,'png_height_px':6267},{'png_width_px':2484,'png_height_px':6267})

if __name__=='__main__':unittest.main()
