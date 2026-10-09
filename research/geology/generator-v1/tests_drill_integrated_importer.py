"""Focused validation for the integrated 11-sheet drilling workbook adapter."""
import sys, tempfile, unittest
from pathlib import Path

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))
from drill_importer import load_drill_workbook
from importer import GeometryInputError

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook=None

NAMES=['钻孔基本信息','回次表','分层-原始记录','分层-地质综合','采样','自定义测试项目','样品测试结果','孔深校正及弯曲度','图签','图面示意','填写说明']


def make_integrated(path):
    wb=Workbook(); wb.remove(wb.active)
    for name in NAMES: wb.create_sheet(name)
    basic=wb['钻孔基本信息']; basic['F4']='SIM-001'; basic['H4']='模拟项目'; basic['F5']=10; basic['H5']='原始沿孔深'; basic['B7']=120.0; basic['D7']='模拟坐标系'; basic['F7']='天顶角自铅垂，方位自北顺时针'; basic['B8']='合成数据，仅用于测试'
    rounds=wb['回次表']; rounds.append([])
    for col,v in enumerate(['序号','回次号','下界记录孔深_m','回次进尺_m','采取块数/原记','岩心长度_m','残留_m','处理后岩心长_m（手填）','回次采取率_pct（原填）','回次采取率_pct（计算）','孔深校正量_m','下界校正孔深_m','≤10cm岩心长度_m','≥10cm岩心长度_m','RQD_pct（原填）','备注','长度核对提示'],1): rounds.cell(5,col).value=v
    rounds['A6']='=IF(B6="","",ROW()-5)'; rounds['B6']='R1'; rounds['C6']=5; rounds['D6']=5; rounds['F6']=4; rounds['H6']=0
    rounds['A7']='=IF(B7="","",ROW()-5)'; rounds['B7']='R2'; rounds['C7']=10; rounds['D7']=5; rounds['F7']=4.8; rounds['H7']=4.5
    layers=wb['分层-原始记录']
    for col,v in enumerate(['序号','层号','起回次号','起位置_m','止回次号','止位置_m','换层累计孔深_m','分层进尺_m','分层岩心长_m','分层采取率_pct（原填）','分层采取率_pct（计算）','岩性名称','平均轴夹角_deg','真厚度_m（原填）','调整/说明','止回次岩心长_m','原表O列未命名值'],1): layers.cell(5,col).value=v
    layers['A6']='=IF(B6="","",ROW()-5)'; layers['B6']='L1'; layers['G6']=5; layers['H6']=5; layers['I6']=3.5; layers['L6']='模拟岩性甲'; layers['N6']=None
    layers['A7']='=IF(B7="","",ROW()-5)'; layers['B7']='L2'; layers['G7']=10; layers['H7']=5; layers['I7']=0; layers['L7']='模拟岩性乙'
    synthesis=wb['分层-地质综合']
    for col,v in enumerate(['序号','起回次号','起位置_m','止回次号','止位置_m','换层孔深_m','分层进尺_m','分层岩心长_m','分层采取率_pct（原填）','平均轴夹角_deg','真厚度_m（原填）'],1): synthesis.cell(5,col).value=v
    samples=wb['采样']
    for col,v in enumerate(['序号','样品编号','起回次号','起位置_m','止回次号','止位置_m','孔深自_m','孔深至_m','进尺_m','岩心长度_m','采取率_pct（原填）','采取率_pct（计算）','重量数值（原单位未明）','重量单位（手填）','原表L列值（含义待确认）','原表M列未命名值/附加备注','原表N列未命名值/附加备注'],1): samples.cell(5,col).value=v
    samples['A6']='=IF(B6="","",ROW()-5)'; samples['B6']='SIM-S1'; samples['G6']=2; samples['H6']=3; samples['I6']=1; samples['J6']=0
    cat=wb['自定义测试项目']
    for col,v in enumerate(['项目代码','项目名称/元素','单位','检出限','分析方法','图中显示','显示顺序','备注'],1): cat.cell(5,col).value=v
    cat.append([]); cat['A6']='X';cat['B6']='模拟项目X';cat['C6']='mg/kg';cat['D6']=0.05;cat['E6']='模拟方法';cat['F6']='否';cat['G6']=2
    cat['A7']='Y';cat['B7']='模拟元素Y';cat['C7']='ppm';cat['D7']=0;cat['E7']='模拟方法2';cat['F7']='是';cat['G7']=1
    res=wb['样品测试结果']
    for col,v in enumerate(['样品编号','项目代码','结果原文','数值结果（可空）','结果标记','单位','检出限','分析方法','检测日期','报告编号','备注'],1): res.cell(5,col).value=v
    res['A6']='SIM-S1';res['B6']='X';res['C6']='<0.05 ';res['E6']='低于检出限';res['F6']='mg/kg';res['G6']=0.05;res['H6']='模拟方法'
    res['A7']='SIM-S1';res['B7']='Y';res['C7']='0';res['D7']=0;res['E7']='定量';res['F7']='ppm'
    depth=wb['孔深校正及弯曲度']
    for col,v in enumerate(['序号','记录孔深_m','校测孔深_m','误差_m','误差率_pct','测量孔深_m','测量天顶角_deg','实测方位角_deg','测量方法','测量仪器'],1): depth.cell(7,col).value=v
    depth['A8']='=IF(B8="","",ROW()-7)';depth['B8']=2.5;depth['C8']=2.6;depth['D8']=0.1;depth['E8']=4;depth['F8']=2.7;depth['G8']=0.8;depth['H8']=90;depth['I8']='模拟测量法';depth['J8']='模拟仪器'
    for c,v in [('A28','应测次数'),('B29',3),('C28','实测次数'),('D29',2),('E28','超差次数'),('F29',1),('G28','应测次数'),('H29',3),('I28','实测次数'),('J29',2),('K28','超差次数'),('L29',0)]: depth[c]=v
    for c,v in [('A30','记录人'),('B30','模拟记录'),('D30','日期'),('E30','2026-10-09'),('G30','检查人'),('H30','模拟检查'),('J30','日期'),('K30','2026-10-09')]: depth[c]=v
    title=wb['图签']
    for c,v in [('B4','模拟项目/单位'),('B5','模拟孔柱状图'),('B6','模拟编制'),('D6','模拟审核'),('B7','模拟制图'),('D7','模拟项目负责'),('B8','模拟单位负责'),('D8','SIM-FIG-1'),('B9','1'),('D9','1000'),('B10','2026-10-09'),('D10','模拟资料')]: title[c]=v
    for r,(w,o) in enumerate([(20,1),(25,2),(20,3),(25,4)],15):
        title[f'A{r}']=f'列{r-14}';title[f'B{r}']=w;title[f'C{r}']=8;title[f'D{r}']=o;title[f'E{r}']='模拟版式'
    title['B19']='90';title['D19']='8'
    wb.save(path);wb.close()


def make_integrated_v3(path):
    make_integrated(path)
    wb=load_workbook(path)
    wb['回次表'].delete_cols(9,1)       # remove old hand-entered recovery rate
    wb['分层-原始记录'].delete_cols(10,1)
    wb['分层-地质综合']['I5']='分层采取率_pct（计算）'
    wb['采样'].delete_cols(11,1)
    wb['样品测试结果'].delete_cols(5,1)  # no result flag in v3
    # Extra synthetic sample and an explicit non-detect preserve raw wording.
    samples=wb['采样'];samples['B7']='SIM-S2';samples['G7']=3;samples['H7']=4;samples['I7']=1
    results=wb['样品测试结果'];results['H5']='检测日期';results['I5']='报告编号';results['A8']='SIM-S2';results['B8']='X';results['C8']='未检出';results['E8']='mg/kg';results['F8']=0.05
    wb.save(path);wb.close()


@unittest.skipIf(Workbook is None,'openpyxl is required')
class IntegratedDrillImporterTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'integrated.xlsx';make_integrated(self.path)
    def tearDown(self):self.tmp.cleanup()
    def codes(self):
        with self.assertRaises(GeometryInputError) as ctx: load_drill_workbook(self.path)
        return {issue['code'] for issue in ctx.exception.issues},ctx.exception.issues
    def test_detects_template_derives_geometry_and_keeps_dynamic_assays_and_provenance(self):
        d=load_drill_workbook(self.path)
        self.assertEqual(d['schema_version'],'drill-integrated-1.0')
        self.assertEqual(d['source']['template_version'],'v2')
        self.assertEqual(d['meta']['hole_id'],'SIM-001')
        self.assertEqual([(x['top_m'],x['bottom_m']) for x in d['layers']],[(0,5),(5,10)])
        self.assertEqual(d['layers'][1]['source']['cells']['顶深_m'],'分层-原始记录!G6')
        self.assertEqual(d['layers'][0]['true_thickness_m_raw'],None)
        self.assertEqual([(x['top_m'],x['bottom_m'],x['core_m']) for x in d['turns']],[(0,5,0),(5,10,4.5)])
        self.assertEqual(d['turns'][1]['source']['derived']['top_m']['source'],'回次表!C6')
        self.assertEqual(d['samples'][0]['assay_raw'],{'X':'<0.05 ','Y':'0'})
        self.assertIsNone(d['samples'][0]['assays']['X'])
        self.assertEqual(d['samples'][0]['assays']['Y'],0)
        self.assertEqual([x['code'] for x in d['project']['analysis_items']],['Y','X'])
        self.assertFalse(next(x for x in d['project']['analysis_items'] if x['code']=='X')['show'])
        self.assertEqual(d['depth_measurements']['records'][0]['recorded_depth_m'],2.5)
        self.assertEqual(d['depth_measurements']['records'][0]['sequence'],1)
        self.assertEqual(d['depth_measurements']['records'][0]['source']['cells']['记录孔深_m'],'孔深校正及弯曲度!B8')
        self.assertEqual(d['title_block']['fields']['项目负责'],'模拟项目负责')
        self.assertEqual(d['title_block']['fields']['单位负责'],'模拟单位负责')
        self.assertEqual(d['title_block']['settings'][0]['width_mm'],20)
        self.assertEqual(d['title_block']['settings_summary']['reference_total_width_mm'],90)
        self.assertEqual(d['basic_info']['fields']['孔口标高_m'],120)
        self.assertEqual(d['turns'][0]['core_m'],0)
        self.assertEqual(d['turns'][0]['recovery_original_percent'],None)
        wb=load_workbook(self.path);wb['回次表']['H6']=None;wb.save(self.path);wb.close()
        d=load_drill_workbook(self.path)
        self.assertIsNone(d['turns'][0]['core_m'])
    def test_dynamic_rows_and_exact_headers(self):
        wb=load_workbook(self.path);cat=wb['自定义测试项目'];cat['A26']='Z';cat['B26']='行26项目';cat['C26']='ppm';cat['F26']=None;wb.save(self.path);wb.close()
        d=load_drill_workbook(self.path)
        self.assertIn('Z',d['project']['analysis_units'])
        wb=load_workbook(self.path);wb['回次表']['C5']='错列名';wb.save(self.path);wb.close()
        codes,_=self.codes()
        self.assertIn('UNSUPPORTED_TABLE_HEADERS',codes)
    def test_cached_formula_columns_are_never_evaluated_or_reported_as_input_errors(self):
        wb=load_workbook(self.path);wb['回次表']['J6']='=1/0';wb.save(self.path);wb.close()
        d=load_drill_workbook(self.path)
        self.assertIsNone(d['turns'][0]['recovery_computed_percent_source'])
        self.assertEqual(d['turns'][0]['recovery_percent'],0)
        self.assertEqual(d['turns'][0]['source_metrics']['computed_recovery_formula'],'=1/0')
    def test_v3_recomputes_recovery_from_raw_lengths_and_preserves_result_text_and_coordinates(self):
        make_integrated_v3(self.path)
        d=load_drill_workbook(self.path)
        self.assertEqual(d['source']['adapter'],'drill-integrated-template-v3')
        self.assertEqual(d['source']['template_version'],'v3')
        self.assertEqual([x['recovery_percent'] for x in d['turns']],[0,90])
        self.assertEqual([x['recovery_percent'] for x in d['layers']],[70,0])
        self.assertIsNone(d['turns'][0]['recovery_original_percent'])
        self.assertEqual(d['turns'][0]['source']['cells']['回次采取率_pct（计算）'],'回次表!I6')
        self.assertEqual(d['layers'][0]['source']['cells']['分层采取率_pct（计算）'],'分层-原始记录!J6')
        self.assertEqual(d['layers'][0]['source']['cells']['岩性名称'],'分层-原始记录!K6')
        self.assertEqual([x['recovery_percent'] for x in d['samples']],[0,None])
        self.assertEqual(d['samples'][0]['assay_raw'],{'X':'<0.05 ','Y':'0'})
        self.assertEqual(d['samples'][0]['assays']['Y'],0)
        self.assertEqual(d['samples'][1]['assay_raw']['X'],'未检出')
        self.assertIsNone(d['samples'][1]['assays']['X'])
        x_assay=next(a for a in d['samples'][0]['assay_records'] if a['code']=='X')
        self.assertIsNone(x_assay['flag'])
        self.assertEqual(x_assay['unit'],'mg/kg')
        self.assertEqual(x_assay['source']['cells']['单位'],'样品测试结果!E6')
        # Rates use hand-entered footage columns D/H/I, never interval depth differences.
        wb=load_workbook(self.path)
        wb['回次表']['H6']=2;wb['回次表']['D6']=None;wb['回次表']['H7']=2;wb['回次表']['D7']=4
        wb['分层-原始记录']['I6']=3.5;wb['分层-原始记录']['H6']=None;wb['分层-原始记录']['I7']=1;wb['分层-原始记录']['H7']=2
        wb['采样']['J6']=0;wb['采样']['I6']=None;wb['采样']['J7']=1;wb['采样']['I7']=2
        wb.save(self.path);wb.close()
        d=load_drill_workbook(self.path)
        self.assertEqual([x['recovery_percent'] for x in d['turns']],[None,50])
        self.assertEqual([x['recovery_percent'] for x in d['layers']],[None,50])
        self.assertEqual([x['recovery_percent'] for x in d['samples']],[None,50])
    def test_formula_only_rows_are_skipped_and_required_formula_is_never_evaluated(self):
        wb=load_workbook(self.path);wb['回次表']['A8']='=1+2';wb['回次表']['J8']='=1/0';wb.save(self.path);wb.close()
        d=load_drill_workbook(self.path)
        self.assertEqual(len(d['turns']),2)
        wb=load_workbook(self.path);wb['回次表']['C6']='=2+3';wb.save(self.path);wb.close()
        codes,issues=self.codes()
        self.assertIn('FORMULA_VALUE',codes)
        self.assertTrue(any('回次表!C6' in issue['cells'] for issue in issues))
    def test_unknown_reference_duplicate_pair_and_unit_conflict_fail(self):
        wb=load_workbook(self.path);ws=wb['样品测试结果'];ws['A7']='SIM-S1';ws['B7']='X';ws['F7']='ppm';ws['A8']='NO-SUCH-SAMPLE';ws['B8']='NO-SUCH-CODE';ws['C8']='<0.01';wb.save(self.path);wb.close()
        codes,_=self.codes()
        self.assertTrue({'DUPLICATE_SAMPLE_ANALYSIS','ANALYSIS_UNIT_CONFLICT','UNKNOWN_SAMPLE','UNKNOWN_ANALYSIS_CODE'}<=codes)

if __name__=='__main__':unittest.main()

