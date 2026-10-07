import math
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from importer import CANONICAL_COLUMNS, GeometryInputError, load_workbook_data


LEGACY_XLS = Path(r"D:\天气地图\地质资料\实测地层剖面登记表.xls")
LEGACY_SHA = "e27d49e7783f6c4414e0dd71477879f335362c9e1daec6489a529492cb2d3c6a"
CANONICAL_PM01 = Path(__file__).resolve().parent / "outputs" / "geology-template-v1" / "PM01规范输入.xlsx"


def write_canonical(path, rows, *, length_unit="m", angle_unit="deg", axis=None, columns=None):
    columns = list(columns or CANONICAL_COLUMNS)
    wb = Workbook()
    project = wb.active
    project.title = "项目"
    project.append(["参数", "值"])
    for key, value in [
        ("项目名称", "合成测试"), ("剖面编号", "T01"),
        ("剖面方位角_deg", axis), ("长度单位", length_unit), ("角度单位", angle_unit),
    ]:
        project.append([key, value])
    segment = wb.create_sheet("测段")
    segment.append(columns)
    for row in rows:
        segment.append([row.get(column) for column in columns])
    wb.save(path)


def base_row(**overrides):
    row = {
        "记录号": "R1", "导线段号": "0-1", "层号": "L1",
        "起读数_m": 0.0, "止读数_m": 10.0, "斜距_m": 10.0,
        "坡角_deg": 0.0, "方位角_deg": 90.0,
        "倾向_deg": 120.0, "倾角_deg": 30.0,
        "岩性名称": "泥岩", "岩性描述": "描述一",
        "样品编号": None, "样品距测段起点_m": None,
        "原表平距_m": 10.0, "原表高差_m": 0.0,
        "原表累计高差_m": 0.0, "原表累计北_m": 0.0,
        "原表累计东_m": 10.0, "实测真厚度_m": None,
    }
    row.update(overrides)
    return row


class LegacyWorkbookTests(unittest.TestCase):
    @unittest.skipUnless(LEGACY_XLS.exists(), "legacy source workbook unavailable")
    def test_pm01_exact_geometry_and_provenance(self):
        data = load_workbook_data(LEGACY_XLS)
        self.assertEqual(data["source"]["sha256"], LEGACY_SHA)
        self.assertEqual(data["summary"]["records"], 43)
        self.assertEqual(len(data["nodes"]), 44)
        self.assertEqual(data["summary"]["intervals"], 26)
        self.assertEqual(data["summary"]["stations"], 20)
        self.assertEqual(len(data["attitudes"]), 27)
        self.assertEqual(data["summary"]["samples"], 14)
        self.assertEqual(data["summary"]["located_samples"], 0)
        self.assertTrue(all(sample["position"] is None for sample in data["samples"]))
        self.assertEqual(sum(len(record["raw"]) for record in data["records"]), 747)
        self.assertAlmostEqual(data["summary"]["endpoint_east_m"], 476.25875170939065, places=12)
        self.assertAlmostEqual(data["summary"]["endpoint_north_m"], 54.25790456876632, places=12)
        self.assertAlmostEqual(data["summary"]["total_vertical_m"], -127.46241495689857, places=12)
        self.assertAlmostEqual(data["summary"]["projected_endpoint_m"], 479.3394609125983, places=12)
        self.assertAlmostEqual(data["summary"]["axis_azimuth_deg"], 83.50058502519266, places=12)
        c12 = [record for record in data["records"] if record["layer_id"] == "C12"]
        self.assertEqual([(r["dip_direction_deg"], r["dip_angle_deg"]) for r in c12], [(300.0, 11.0), (300.0, 11.0)])
        self.assertIsNone(c12[0]["true_thickness_m"])
        first = data["records"][0]
        self.assertEqual(first["source_cells"]["start_reading_m"], "C6")
        self.assertEqual(first["source_cells"]["end_reading_m"], "D6")
        self.assertEqual(first["raw"]["C6"], 0.0)
        self.assertEqual(first["raw"]["D6"], 43.4)
        self.assertTrue(any(issue["code"] == "CACHE_MISMATCH" for issue in data["issues"]))


class CanonicalWorkbookTests(unittest.TestCase):
    def run_book(self, rows, **kwargs):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            write_canonical(path, rows, **kwargs)
            return load_workbook_data(path)

    @unittest.skipUnless(CANONICAL_PM01.exists(), "canonical PM01 workbook unavailable")
    def test_artifact_tool_xlsx_without_dimension_roundtrips(self):
        data = load_workbook_data(CANONICAL_PM01)
        self.assertEqual((data["summary"]["records"], data["summary"]["stations"]), (43, 20))
        self.assertAlmostEqual(data["summary"]["endpoint_east_m"], 476.25875170939065, places=12)
        self.assertAlmostEqual(data["summary"]["endpoint_north_m"], 54.25790456876632, places=12)
        self.assertAlmostEqual(data["summary"]["total_vertical_m"], -127.46241495689857, places=12)

    def test_reordered_columns_dynamic_counts_and_samples(self):
        rows = [
            base_row(**{"样品编号": "S1", "样品距测段起点_m": 2.5, "实测真厚度_m": 1.25}),
            base_row(**{"记录号": "R2", "导线段号": "0-1", "起读数_m": 10.0,
                        "止读数_m": 15.0, "斜距_m": 5.0, "层号": "L1",
                        "岩性名称": None, "岩性描述": None, "样品编号": "S2",
                        "原表平距_m": None, "原表累计东_m": None}),
            base_row(**{"记录号": "R3", "导线段号": "1-2", "层号": "L2",
                        "起读数_m": 0.0, "止读数_m": 4.0, "斜距_m": 4.0,
                        "方位角_deg": 0.0, "倾向_deg": None, "倾角_deg": None,
                        "岩性名称": "灰岩", "岩性描述": "描述二", "原表平距_m": None,
                        "原表累计东_m": None}),
        ]
        columns = list(reversed(CANONICAL_COLUMNS))
        data = self.run_book(rows, columns=columns)
        self.assertEqual(data["summary"]["records"], 3)
        self.assertEqual(data["summary"]["intervals"], 2)
        self.assertEqual(data["summary"]["stations"], 3)
        self.assertEqual(len(data["nodes"]), 4)
        self.assertEqual(data["records"][1]["lithology_name"], "泥岩")
        self.assertEqual(data["records"][1]["description"], "描述一")
        self.assertEqual(data["samples"][0]["location_status"], "explicit_offset")
        self.assertAlmostEqual(data["samples"][0]["position"]["east_m"], 2.5)
        self.assertEqual(data["samples"][1]["location_status"], "missing")
        self.assertEqual(data["records"][0]["true_thickness_m"], 1.25)

    def test_explicit_lithology_survives_blank_description_and_then_inherits(self):
        rows = [
            base_row(**{"岩性名称": "凝灰岩", "岩性描述": None}),
            base_row(**{"记录号": "R2", "岩性名称": None, "岩性描述": None,
                        "起读数_m": 10.0, "止读数_m": 20.0}),
        ]
        data = self.run_book(rows)
        self.assertEqual([r["lithology_name"] for r in data["records"]], ["凝灰岩", "凝灰岩"])
        self.assertEqual([r["description"] for r in data["records"]], ["", ""])

    def test_repeated_explicit_leg_is_one_station_span(self):
        rows = [base_row(), base_row(**{"记录号": "R2", "起读数_m": 10, "止读数_m": 20})]
        data = self.run_book(rows)
        self.assertEqual([(s["id"], s["node"]) for s in data["stations"]], [("0", 0), ("1", 2)])

    def test_noncontiguous_same_layer_makes_three_intervals(self):
        rows = [
            base_row(),
            base_row(**{"记录号": "R2", "导线段号": "1-2", "层号": "L2", "岩性名称": "灰岩"}),
            base_row(**{"记录号": "R3", "导线段号": "2-3", "层号": "L1", "岩性名称": "泥岩"}),
        ]
        self.assertEqual(self.run_book(rows)["summary"]["intervals"], 3)

    def assert_error_code(self, rows, code, **kwargs):
        with self.assertRaises(GeometryInputError) as caught:
            self.run_book(rows, **kwargs)
        self.assertIn(code, [issue["code"] for issue in caught.exception.issues])

    def test_invalid_unit_rejected(self):
        self.assert_error_code([base_row()], "INVALID_LENGTH_UNIT", length_unit="ft")

    def test_missing_geometry_rejected(self):
        self.assert_error_code([base_row(**{"斜距_m": None})], "MISSING_GEOMETRY")

    def test_boolean_geometry_rejected(self):
        self.assert_error_code([base_row(**{"斜距_m": True})], "INVALID_NUMBER")

    def test_duplicate_record_id_rejected(self):
        self.assert_error_code([base_row(), base_row()], "DUPLICATE_RECORD_ID")

    def test_conflicting_lithology_in_one_layer_rejected(self):
        rows = [base_row(), base_row(**{"记录号": "R2", "岩性名称": "灰岩"})]
        self.assert_error_code(rows, "LITHOLOGY_CONFLICT")

    def test_disconnected_parseable_legs_rejected(self):
        rows = [base_row(), base_row(**{"记录号": "R2", "导线段号": "3-4", "层号": "L2"})]
        self.assert_error_code(rows, "DISCONNECTED_STATIONS")

    def test_reading_length_conflict_warns_without_overwrite(self):
        data = self.run_book([base_row(**{"止读数_m": 12.0})])
        self.assertEqual(data["records"][0]["length_m"], 10.0)
        self.assertIn("READING_LENGTH_MISMATCH", [issue["code"] for issue in data["issues"]])


if __name__ == "__main__":
    unittest.main(verbosity=2)
