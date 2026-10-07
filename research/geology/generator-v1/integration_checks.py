"""End-to-end acceptance checks for the geology generator v1.

The checks use only the original legacy workbook, the newly generated canonical
workbook, importer.py, renderer.py, and temporary workbook fixtures.  They do
not read prior drawings, legacy JSON, photographs, or independent baselines.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import sys
import tempfile
import traceback
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from importer import GeometryInputError, load_workbook_data
from renderer import render


ROOT = Path(__file__).resolve().parent
LEGACY_XLS = Path(r"D:\天气地图\地质资料\实测地层剖面登记表.xls")
CANONICAL_XLSX = ROOT / "outputs" / "geology-template-v1" / "PM01规范输入.xlsx"
MATERIALS_JSON = ROOT / "templates" / "materials.json"
RESULT_JSON = ROOT / "logs" / "integration-checks.json"
FLOAT_TOLERANCE = 1e-12


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def semantic_payload(data: dict[str, Any]) -> dict[str, Any]:
    record_keys = (
        "id", "leg_id", "layer_id", "length_m", "slope_deg", "azimuth_deg",
        "dip_direction_deg", "dip_angle_deg", "lithology_name", "description",
        "start_node", "end_node", "computed", "true_thickness_m",
    )
    interval_keys = (
        "id", "layer_id", "start_node", "end_node", "record_ids",
        "lithology_name", "description",
    )
    station_keys = ("id", "node")
    return {
        "project": data.get("project"),
        "settings": {
            "axis_azimuth_deg": data.get("settings", {}).get("axis_azimuth_deg"),
            "coordinate_system": data.get("settings", {}).get("coordinate_system"),
            "vertical_exaggeration": data.get("settings", {}).get("vertical_exaggeration"),
        },
        "records": [{key: record.get(key) for key in record_keys} for record in data.get("records", [])],
        "nodes": data.get("nodes", []),
        "intervals": [{key: item.get(key) for key in interval_keys} for item in data.get("intervals", [])],
        "stations": [{key: item.get(key) for key in station_keys} for item in data.get("stations", [])],
        "attitudes": data.get("attitudes", []),
        "samples": [
            {key: item.get(key) for key in ("id", "record_id", "layer_id", "position", "location_status")}
            for item in data.get("samples", [])
        ],
        "summary": data.get("summary", {}),
    }


def compare_recursive(left: Any, right: Any, path: str = "root", tolerance: float = FLOAT_TOLERANCE) -> float:
    """Compare semantic structures and return the greatest numeric delta."""
    if isinstance(left, bool) or isinstance(right, bool):
        require(left is right, f"{path}: boolean mismatch {left!r} != {right!r}")
        return 0.0
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        require(math.isfinite(float(left)) and math.isfinite(float(right)), f"{path}: non-finite number")
        delta = abs(float(left) - float(right))
        require(delta <= tolerance, f"{path}: numeric delta {delta:.17g} exceeds {tolerance:g}")
        return delta
    require(type(left) is type(right), f"{path}: type mismatch {type(left).__name__} != {type(right).__name__}")
    if isinstance(left, dict):
        require(left.keys() == right.keys(), f"{path}: key mismatch {sorted(left)} != {sorted(right)}")
        return max((compare_recursive(left[key], right[key], f"{path}.{key}", tolerance) for key in left), default=0.0)
    if isinstance(left, list):
        require(len(left) == len(right), f"{path}: list length {len(left)} != {len(right)}")
        return max((compare_recursive(a, b, f"{path}[{index}]", tolerance) for index, (a, b) in enumerate(zip(left, right))), default=0.0)
    require(left == right, f"{path}: value mismatch {left!r} != {right!r}")
    return 0.0


def sheet_columns(sheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in sheet[1] if cell.value is not None}


def project_row(sheet, parameter: str) -> int:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == parameter:
            return row
    raise AssertionError(f"项目参数不存在: {parameter}")


def edit_copy(source: Path, destination: Path, editor: Callable[[Any], None]) -> None:
    shutil.copy2(source, destination)
    workbook = load_workbook(destination)
    editor(workbook)
    workbook.save(destination)
    workbook.close()


def polygon_pattern(svg_text: str, interval_id: str) -> str | None:
    match = re.search(
        rf'<polygon\b[^>]*fill="url\(#([^\)]+)\)"[^>]*data-interval="{re.escape(interval_id)}"',
        svg_text,
    )
    return match.group(1) if match else None


def issue_codes(error: GeometryInputError) -> list[str]:
    return [str(issue.get("code")) for issue in error.issues]


def main() -> int:
    started = datetime.now(timezone.utc).isoformat()
    RESULT_JSON.parent.mkdir(parents=True, exist_ok=True)
    require(LEGACY_XLS.exists(), f"缺少原始 XLS: {LEGACY_XLS}")
    require(CANONICAL_XLSX.exists(), f"缺少规范 XLSX: {CANONICAL_XLSX}")
    source_sha_before = sha256(LEGACY_XLS)
    canonical_sha_before = sha256(CANONICAL_XLSX)
    legacy = load_workbook_data(LEGACY_XLS)
    canonical = load_workbook_data(CANONICAL_XLSX)
    configured_materials = list(json.loads(MATERIALS_JSON.read_text(encoding="utf-8"))["materials"])
    results: list[dict[str, Any]] = []

    def check(name: str, callback: Callable[[], dict[str, Any] | None]) -> None:
        try:
            detail = callback() or {}
            results.append({"name": name, "status": "pass", "detail": detail})
        except Exception as exc:  # keep all check evidence in one JSON report
            results.append({
                "name": name,
                "status": "fail",
                "error_type": type(exc).__name__,
                "error": str(exc),
                "traceback": traceback.format_exc(limit=8),
            })

    with tempfile.TemporaryDirectory(prefix="geology-v1-integration-") as temporary:
        temp = Path(temporary)

        def check_roundtrip() -> dict[str, Any]:
            delta = compare_recursive(semantic_payload(legacy), semantic_payload(canonical))
            require(legacy["summary"]["records"] == 43, "legacy record count changed")
            require(canonical["summary"]["stations"] == 20, "canonical station count is not 20")
            return {
                "records": canonical["summary"]["records"],
                "intervals": canonical["summary"]["intervals"],
                "stations": canonical["summary"]["stations"],
                "attitudes": len(canonical["attitudes"]),
                "samples": len(canonical["samples"]),
                "max_numeric_delta": delta,
                "tolerance": FLOAT_TOLERANCE,
            }

        check("legacy_to_canonical_semantic_roundtrip", check_roundtrip)

        def check_length_recalculation() -> dict[str, Any]:
            target_index = min(10, len(canonical["records"]) - 1)
            record = canonical["records"][target_index]
            changed_length = float(record["length_m"]) + 1.25
            fixture = temp / "length-change.xlsx"

            def editor(workbook) -> None:
                sheet = workbook["测段"]
                columns = sheet_columns(sheet)
                sheet.cell(target_index + 2, columns["斜距_m"]).value = changed_length

            edit_copy(CANONICAL_XLSX, fixture, editor)
            changed = load_workbook_data(fixture)
            require(changed["records"][target_index]["length_m"] == changed_length, "changed length was not imported")
            for index in range(target_index + 1):
                compare_recursive(canonical["nodes"][index], changed["nodes"][index], f"unchanged_node[{index}]", FLOAT_TOLERANCE)
            slope = math.radians(float(record["slope_deg"]))
            azimuth = math.radians(float(record["azimuth_deg"]))
            delta_length = changed_length - float(record["length_m"])
            delta_horizontal = delta_length * math.cos(slope)
            expected = {
                "east_m": delta_horizontal * math.sin(azimuth),
                "north_m": delta_horizontal * math.cos(azimuth),
                "z_m": delta_length * math.sin(slope),
                "chainage_m": delta_horizontal,
                "slant_chainage_m": delta_length,
            }
            axis = math.radians(float(changed["settings"]["axis_azimuth_deg"]))
            expected["x_m"] = expected["east_m"] * math.sin(axis) + expected["north_m"] * math.cos(axis)
            expected["offset_m"] = expected["north_m"] * math.sin(axis) - expected["east_m"] * math.cos(axis)
            max_delta = 0.0
            for node_index in range(target_index + 1, len(changed["nodes"])):
                for field, expected_shift in expected.items():
                    actual_shift = float(changed["nodes"][node_index][field]) - float(canonical["nodes"][node_index][field])
                    delta = abs(actual_shift - expected_shift)
                    max_delta = max(max_delta, delta)
                    require(delta <= FLOAT_TOLERANCE, f"node {node_index} {field} was not fully recalculated")
            require(any(abs(value) > 1e-9 for value in expected.values()), "length edit produced no geometric change")
            return {"record_id": record["id"], "delta_length_m": delta_length, "successor_nodes_checked": len(changed["nodes"]) - target_index - 1, "max_delta": max_delta}

        check("length_change_recalculates_all_successor_nodes", check_length_recalculation)

        layer_counts = Counter(record["layer_id"] for record in canonical["records"])
        unique_index = next(index for index, record in enumerate(canonical["records"]) if layer_counts[record["layer_id"]] == 1)
        unique_record = canonical["records"][unique_index]
        original_material = unique_record["lithology_name"]
        alternate_material = next(name for name in configured_materials if name != original_material)

        def make_lithology_fixture(name: str, filename: str) -> Path:
            fixture = temp / filename

            def editor(workbook) -> None:
                sheet = workbook["测段"]
                columns = sheet_columns(sheet)
                sheet.cell(unique_index + 2, columns["岩性名称"]).value = name

            edit_copy(CANONICAL_XLSX, fixture, editor)
            return fixture

        def check_lithology_only() -> dict[str, Any]:
            fixture = make_lithology_fixture(alternate_material, "lithology-change.xlsx")
            changed = load_workbook_data(fixture)
            compare_recursive(canonical["nodes"], changed["nodes"], "nodes", FLOAT_TOLERANCE)
            target_interval = next(item for item in changed["intervals"] if unique_record["id"] in item["record_ids"])
            require(target_interval["lithology_name"] == alternate_material, "new material did not reach target interval")
            old_output, new_output = temp / "lith-old", temp / "lith-new"
            old_paths = render(canonical, old_output, {"create_preview_png": False})
            new_paths = render(changed, new_output, {"create_preview_png": False})
            old_svg = Path(old_paths["drawing_svg"]).read_text(encoding="utf-8")
            new_svg = Path(new_paths["drawing_svg"]).read_text(encoding="utf-8")
            old_pattern = polygon_pattern(old_svg, target_interval["id"])
            new_pattern = polygon_pattern(new_svg, target_interval["id"])
            require(old_pattern is not None and new_pattern is not None, "target interval polygon not found")
            require(old_pattern != new_pattern, f"pattern did not change: {old_pattern}")
            return {"record_id": unique_record["id"], "interval_id": target_interval["id"], "from": original_material, "to": alternate_material, "old_pattern": old_pattern, "new_pattern": new_pattern}

        check("lithology_only_changes_pattern_not_nodes", check_lithology_only)

        def check_unknown_pending() -> dict[str, Any]:
            unknown = "未配置测试岩性"
            fixture = make_lithology_fixture(unknown, "unknown-lithology.xlsx")
            changed = load_workbook_data(fixture)
            output = temp / "unknown-output"
            paths = render(changed, output, {"create_preview_png": False})
            audit = json.loads(Path(paths["layout_audit_json"]).read_text(encoding="utf-8"))
            target_interval = next(item for item in changed["intervals"] if unique_record["id"] in item["record_ids"])
            svg = Path(paths["drawing_svg"]).read_text(encoding="utf-8")
            require(audit["materials"]["pending_names"] == [unknown], f"unexpected pending names: {audit['materials']['pending_names']}")
            require(audit["materials"]["fuzzy_matching"] is False, "fuzzy matching was enabled")
            require(unknown not in audit["materials"]["configured_names"], "unknown name was classified")
            pending_pattern = next(item for item in audit["pattern_audit"] if item["interval_id"] == target_interval["id"])
            require(pending_pattern["template_pattern_id"] == "mat_pending", "unknown interval did not use pending template")
            require(
                polygon_pattern(svg, target_interval["id"]) == pending_pattern["derived_pattern_id"],
                "unknown interval polygon did not use its pending derived pattern",
            )
            return {
                "unknown_name": unknown,
                "interval_id": target_interval["id"],
                "template_pattern": "mat_pending",
                "derived_pattern": pending_pattern["derived_pattern_id"],
            }

        check("unknown_lithology_remains_pending", check_unknown_pending)

        def check_small_dataset() -> dict[str, Any]:
            fixture = temp / "small.xlsx"

            def editor(workbook) -> None:
                sheet = workbook["测段"]
                if sheet.max_row > 3:
                    sheet.delete_rows(4, sheet.max_row - 3)

            edit_copy(CANONICAL_XLSX, fixture, editor)
            data = load_workbook_data(fixture)
            require((len(data["records"]), len(data["nodes"]), len(data["intervals"])) == (2, 3, 1), "small dataset counts are wrong")
            paths = render(data, temp / "small-output", {"create_preview_png": False})
            for key in ("drawing_svg", "layout_audit_json", "appendix_html"):
                require(Path(paths[key]).exists() and Path(paths[key]).stat().st_size > 0, f"small dataset missing {key}")
            return {"records": 2, "nodes": 3, "intervals": 1}

        check("small_dynamic_dataset_generates", check_small_dataset)

        def expect_import_error(fixture: Path, expected_code: str) -> list[str]:
            try:
                load_workbook_data(fixture)
            except GeometryInputError as error:
                codes = issue_codes(error)
                require(expected_code in codes, f"expected {expected_code}, got {codes}")
                return sorted(set(codes))
            raise AssertionError(f"fixture did not reject with {expected_code}")

        def check_missing_geometry() -> dict[str, Any]:
            fixture = temp / "missing-geometry.xlsx"

            def editor(workbook) -> None:
                sheet = workbook["测段"]
                sheet.cell(2, sheet_columns(sheet)["斜距_m"]).value = None

            edit_copy(CANONICAL_XLSX, fixture, editor)
            return {"issue_codes": expect_import_error(fixture, "MISSING_GEOMETRY")}

        check("missing_geometry_rejected", check_missing_geometry)

        def check_wrong_unit() -> dict[str, Any]:
            fixture = temp / "wrong-unit.xlsx"

            def editor(workbook) -> None:
                sheet = workbook["项目"]
                sheet.cell(project_row(sheet, "长度单位"), 2).value = "cm"

            edit_copy(CANONICAL_XLSX, fixture, editor)
            return {"issue_codes": expect_import_error(fixture, "INVALID_LENGTH_UNIT")}

        check("wrong_unit_rejected", check_wrong_unit)

        deterministic_output_1 = temp / "deterministic-1"
        deterministic_output_2 = temp / "deterministic-2"

        def check_deterministic_svg() -> dict[str, Any]:
            first = render(canonical, deterministic_output_1, {"create_preview_png": False})
            second = render(canonical, deterministic_output_2, {"create_preview_png": False})
            first_sha = sha256(Path(first["drawing_svg"]))
            second_sha = sha256(Path(second["drawing_svg"]))
            require(first_sha == second_sha, "same input produced different SVG bytes")
            return {"svg_sha256": first_sha}

        check("same_input_svg_is_deterministic", check_deterministic_svg)

        def check_missing_samples_not_plotted() -> dict[str, Any]:
            paths = render(canonical, temp / "missing-samples", {"create_preview_png": False})
            audit = json.loads(Path(paths["layout_audit_json"]).read_text(encoding="utf-8"))
            svg = Path(paths["drawing_svg"]).read_text(encoding="utf-8")
            missing_ids = [sample["id"] for sample in canonical["samples"] if sample["location_status"] == "missing"]
            require(missing_ids, "fixture has no missing-location samples")
            require(audit["samples"]["missing_location"] == len(missing_ids), "audit missing sample count mismatch")
            require(audit["samples"]["missing_location_plotted"] is False, "audit says missing samples were plotted")
            require('data-sample="' not in svg, "missing-location sample marker appeared in SVG")
            return {"missing_samples": len(missing_ids), "plotted_markers": 0}

        check("missing_sample_positions_are_not_plotted", check_missing_samples_not_plotted)

        def check_explicit_sample_interpolation() -> dict[str, Any]:
            sample = canonical["samples"][0]
            record_index = next(index for index, record in enumerate(canonical["records"]) if record["id"] == sample["record_id"])
            record = canonical["records"][record_index]
            distance = float(record["length_m"]) * 0.375
            fixture = temp / "located-sample.xlsx"

            def editor(workbook) -> None:
                sheet = workbook["测段"]
                columns = sheet_columns(sheet)
                sheet.cell(record_index + 2, columns["样品距测段起点_m"]).value = distance

            edit_copy(CANONICAL_XLSX, fixture, editor)
            changed = load_workbook_data(fixture)
            located = next(item for item in changed["samples"] if item["id"] == sample["id"])
            require(located["location_status"] == "explicit_offset", "explicit sample offset was not accepted")
            start = changed["nodes"][record["start_node"]]
            end = changed["nodes"][record["end_node"]]
            max_delta = 0.0
            for field in ("east_m", "north_m", "z_m", "x_m", "offset_m"):
                expected = float(start[field]) + (float(end[field]) - float(start[field])) * 0.375
                delta = abs(float(located["position"][field]) - expected)
                max_delta = max(max_delta, delta)
                require(delta <= FLOAT_TOLERANCE, f"sample {field} interpolation mismatch")
            paths = render(changed, temp / "located-sample-output", {"create_preview_png": False})
            svg = Path(paths["drawing_svg"]).read_text(encoding="utf-8")
            require(f'data-sample="{sample["id"]}"' in svg, "explicitly located sample marker missing")
            return {"sample_id": sample["id"], "fraction": 0.375, "max_numeric_delta": max_delta}

        check("explicit_sample_offset_interpolates_and_plots", check_explicit_sample_interpolation)

    source_sha_after = sha256(LEGACY_XLS)
    canonical_sha_after = sha256(CANONICAL_XLSX)

    def check_source_immutability() -> dict[str, Any]:
        require(source_sha_before == source_sha_after, "original XLS SHA changed during checks")
        require(canonical_sha_before == canonical_sha_after, "final canonical XLSX was modified by temporary checks")
        return {"legacy_sha256": source_sha_after, "canonical_sha256": canonical_sha_after}

    check("source_and_final_workbook_unchanged", check_source_immutability)
    passed = sum(item["status"] == "pass" for item in results)
    report = {
        "schema_version": "1.0",
        "started_at_utc": started,
        "finished_at_utc": datetime.now(timezone.utc).isoformat(),
        "inputs": {
            "legacy_xls": str(LEGACY_XLS),
            "canonical_xlsx": str(CANONICAL_XLSX),
            "legacy_sha256_before": source_sha_before,
            "canonical_sha256_before": canonical_sha_before,
        },
        "temporary_fixtures_only": True,
        "prior_outputs_used_as_input": False,
        "summary": {"total": len(results), "passed": passed, "failed": len(results) - passed, "status": "pass" if passed == len(results) else "fail"},
        "checks": results,
    }
    RESULT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(RESULT_JSON), **report["summary"]}, ensure_ascii=False))
    return 0 if report["summary"]["status"] == "pass" else 1


if __name__ == "__main__":
    sys.exit(main())
