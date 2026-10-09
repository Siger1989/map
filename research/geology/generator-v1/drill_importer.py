"""Read and validate drill-1.0 workbooks without inferring missing geology."""
from __future__ import annotations

import hashlib
import math
from pathlib import Path
from typing import Any

from importer import GeometryInputError, _cell, _is_blank, _issue, _number, _text
from drill_integrated_importer import INTEGRATED_SHEETS, load_integrated_drill_workbook

PROJECT_REQUIRED = ["模板类型", "模板版本", "项目名称", "钻孔编号", "长度单位", "孔深基准", "终孔深度_m", "Au单位", "Pb单位", "Zn单位"]
SHEETS = {
    "分层": ["层号", "顶深_m", "底深_m", "岩性名称", "岩性描述", "岩心长_m", "花纹代码"],
    "回次": ["回次号", "顶深_m", "底深_m", "岩心长_m"],
    "样品": ["样品编号", "顶深_m", "底深_m", "岩心长_m", "Au", "Pb", "Zn"],
    "孔径": ["孔深_m", "孔径_mm"],
}


def _source(sheet: str, row: int, cells: dict[str, str]) -> dict[str, Any]:
    return {"sheet": sheet, "row": row, "cells": cells}


def _hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _sheet_rows(ws, required: list[str], issues: list[dict[str, Any]]):
    headers = [_text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    duplicates = sorted({h for h in required if h and headers.count(h) > 1})
    missing = [x for x in required if x not in headers]
    if duplicates:
        issues.append(_issue("error", "DUPLICATE_COLUMN", f"{ws.title}表存在重复列名：{', '.join(duplicates)}", [f"{ws.title}!1:1"]))
    if missing:
        issues.append(_issue("error", "MISSING_COLUMN", f"{ws.title}表缺少列：{', '.join(missing)}", [f"{ws.title}!1:1"]))
    cols = {name: headers.index(name) + 1 for name in required if name in headers}
    rows = []
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        vals = {name: (cells[col - 1].value if col <= len(cells) else None) for name, col in cols.items()}
        if all(_is_blank(v) for v in vals.values()):
            continue
        formulas = [f"{ws.title}!{_cell(c-1,rn)}" for name,c in cols.items()
                    if isinstance(vals.get(name), str) and vals[name].startswith("=")]
        if formulas:
            issues.append(_issue("error", "FORMULA_VALUE", "数据单元格含公式；请粘贴计算后的数值", formulas))
        rows.append((rn, vals, {name: f"{ws.title}!{_cell(col-1,rn)}" for name,col in cols.items()}))
    return rows


def _load_drill_workbook_open(path: str | Path, wb) -> dict[str, Any]:
    path = Path(path)
    issues: list[dict[str, Any]] = []
    if "项目" not in wb.sheetnames:
        wb.close()
        raise GeometryInputError([_issue("error", "MISSING_SHEET", "规范工作簿必须包含“项目”工作表", [])])
    ws = wb["项目"]
    if _text(ws["A1"].value) != "参数" or _text(ws["B1"].value) != "值":
        issues.append(_issue("error", "INVALID_PROJECT_HEADER", "项目表 A1/B1 必须为“参数”/“值”", ["项目!A1", "项目!B1"]))
    params, pcells = {}, {}
    for rn, row in enumerate(ws.iter_rows(min_row=2), 2):
        key = _text(row[0].value if row else None)
        if key:
            value = row[1].value if len(row) > 1 else None
            if key in params:
                issues.append(_issue("error", "DUPLICATE_PROJECT_PARAMETER", f"项目参数重复：{key}", [pcells[key], f"项目!B{rn}"]))
            if isinstance(value, str) and value.startswith("="):
                issues.append(_issue("error", "FORMULA_VALUE", "项目参数含公式；请粘贴计算后的值", [f"项目!B{rn}"]))
            params[key] = row[1].value if len(row) > 1 else None
            pcells[key] = f"项目!B{rn}"
    for key in PROJECT_REQUIRED:
        if key not in params:
            issues.append(_issue("error", "MISSING_PROJECT_PARAMETER", f"项目表缺少参数：{key}", ["项目!A:A"]))
    if _text(params.get("模板类型")) != "钻孔柱状图":
        issues.append(_issue("error", "INVALID_TEMPLATE_TYPE", "模板类型必须为“钻孔柱状图”", [pcells.get("模板类型", "项目!B:B")]))
    version=params.get("模板版本")
    if not ((isinstance(version,(int,float)) and not isinstance(version,bool) and float(version)==1.0) or _text(version)=="1.0"):
        issues.append(_issue("error", "INVALID_TEMPLATE_VERSION", "模板版本必须为 1.0", [pcells.get("模板版本", "项目!B:B")]))
    if _text(params.get("长度单位")) != "m":
        issues.append(_issue("error", "INVALID_LENGTH_UNIT", "长度单位必须为 m", [pcells.get("长度单位", "项目!B:B")]))
    basis = _text(params.get("孔深基准"))
    if basis not in ("原始沿孔深", "校正沿孔深"):
        issues.append(_issue("error", "INVALID_DEPTH_BASIS", "孔深基准必须明确为原始沿孔深或校正沿孔深", [pcells.get("孔深基准", "项目!B:B")]))
    hole_id = _text(params.get("钻孔编号"))
    if not hole_id:
        issues.append(_issue("error", "MISSING_HOLE_ID", "钻孔编号必填", [pcells.get("钻孔编号", "项目!B:B")]))
    endpoint = _number(params.get("终孔深度_m"), pcells.get("终孔深度_m", "项目!B:B"), issues, required=True)
    if endpoint is not None and endpoint <= 0:
        issues.append(_issue("error", "INVALID_ENDPOINT", "终孔深度必须大于 0", [pcells["终孔深度_m"]]))
    recordsets = {}
    for name, required in SHEETS.items():
        if name in wb.sheetnames:
            recordsets[name] = _sheet_rows(wb[name], required, issues)
        else:
            recordsets[name] = []
    if "分层" not in wb.sheetnames:
        issues.append(_issue("error", "MISSING_LAYER_SHEET", "必需的“分层”工作表缺失", []))
    layers, turns, samples, structures = [], [], [], []
    def num(v, cell, req=False):
        return _number(v, cell, issues, required=req)
    seen = {"分层": set(), "回次": set(), "样品": set()}
    for rn, v, c in recordsets["分层"]:
        ident = _text(v.get("层号")); top = num(v.get("顶深_m"), c.get("顶深_m", f"分层!row{rn}"), True); bot = num(v.get("底深_m"), c.get("底深_m", f"分层!row{rn}"), True)
        if not ident: issues.append(_issue("error", "MISSING_ID", "层号不能为空", [c.get("层号", f"分层!A{rn}")]))
        elif ident in seen["分层"]: issues.append(_issue("error", "DUPLICATE_ID", f"层号重复：{ident}", [c["层号"]]))
        seen["分层"].add(ident)
        core = num(v.get("岩心长_m"), c.get("岩心长_m", ""))
        if core is not None and core < 0: issues.append(_issue("error", "NEGATIVE_CORE_LENGTH", "岩心长不能为负数", [c.get("岩心长_m", "")]))
        if core is not None and top is not None and bot is not None and core > bot-top+1e-6: issues.append(_issue("error", "CORE_LENGTH_EXCEEDS_INTERVAL", "岩心长不能大于层段沿孔长度", [c.get("岩心长_m", ""), c.get("顶深_m", ""), c.get("底深_m", "")]))
        layers.append({"id": ident, "top_m": top, "bottom_m": bot, "thickness_m": bot-top if top is not None and bot is not None else None, "core_m": core, "recovery_percent": core/(bot-top)*100 if core is not None and top is not None and bot is not None and bot>top else None, "description": _text(v.get("岩性描述")), "lithology_name": _text(v.get("岩性名称")), "material_code": _text(v.get("花纹代码")), "pattern_id": _text(v.get("花纹代码")) or None, "source": _source("分层",rn,c)})
        if not _text(v.get("岩性名称")): issues.append(_issue("warning", "MISSING_LITHOLOGY", "岩性名称为空，保留空白待补", [c.get("岩性名称", "")]))
    if not layers:
        issues.append(_issue("error", "EMPTY_LAYER_SHEET", "分层表至少需要一条有效分层记录", ["分层!2:2"]))
    def check_ranges(items, label):
        valid = [x for x in items if x["top_m"] is not None and x["bottom_m"] is not None]
        # Row order is meaningful source evidence. Reject out-of-order rows instead of sorting away the issue.
        ordered = valid
        cursor=0.0
        if not ordered: return
        for x in ordered:
            if x["bottom_m"] <= x["top_m"]: issues.append(_issue("error", "INVALID_INTERVAL", f"{label}底深必须大于顶深", [x["source"]["cells"].get("顶深_m", ""),x["source"]["cells"].get("底深_m", "")]))
            if x["top_m"] < cursor-1e-6: issues.append(_issue("error", "INTERVAL_OVERLAP", f"{label}区间重叠或倒序", [x["source"]["cells"].get("顶深_m", "")]))
            elif x["top_m"] > cursor+1e-6: issues.append(_issue("error", "INTERVAL_GAP", f"{label}区间存在缺段：{cursor:g}–{x['top_m']:g} m", [x["source"]["cells"].get("顶深_m", "")]))
            cursor=max(cursor,x["bottom_m"])
            if x["top_m"] < -1e-6 or endpoint is not None and x["bottom_m"] > endpoint+1e-6:
                issues.append(_issue("error", "INTERVAL_OUT_OF_RANGE", f"{label}区间超出0至终孔深度", [x["source"]["cells"].get("顶深_m", ""),x["source"]["cells"].get("底深_m", "")]))
        if ordered[0]["top_m"]>1e-6 or endpoint is not None and abs(cursor-endpoint)>1e-6:
            issues.append(_issue("error", "INCOMPLETE_COVERAGE", f"{label}必须连续覆盖0至终孔深度", [ordered[0]["source"]["cells"].get("顶深_m", ""),ordered[-1]["source"]["cells"].get("底深_m", "")]))
    check_ranges(layers,"分层")
    for rn,v,c in recordsets["回次"]:
        ident=_text(v.get("回次号")); top=num(v.get("顶深_m"),c.get("顶深_m",""),True); bot=num(v.get("底深_m"),c.get("底深_m",""),True); core=num(v.get("岩心长_m"),c.get("岩心长_m",""),True)
        if not ident: issues.append(_issue("error","MISSING_ID","回次号不能为空",[c.get("回次号","")]))
        elif ident in seen["回次"]: issues.append(_issue("error","DUPLICATE_ID",f"回次号重复：{ident}",[c["回次号"]]))
        seen["回次"].add(ident)
        if top is not None and bot is not None and bot>top and core is not None and core>bot-top+1e-6: issues.append(_issue("error","CORE_LENGTH_EXCEEDS_INTERVAL","岩心长不能大于回次段长",[c.get("岩心长_m","")]))
        if core is not None and core<0: issues.append(_issue("error","NEGATIVE_CORE_LENGTH","岩心长不能为负数",[c.get("岩心长_m","")]))
        turns.append({"id":ident,"top_m":top,"bottom_m":bot,"advance_m":bot-top if top is not None and bot is not None else None,"core_m":core,"recovery_percent":core/(bot-top)*100 if core is not None and top is not None and bot is not None and bot>top else None,"source":_source("回次",rn,c)})
    if turns: check_ranges(turns,"回次")
    units={e:_text(params.get(f"{e}单位")) for e in ("Au","Pb","Zn")}; assay_data={e:[] for e in units}
    for rn,v,c in recordsets["样品"]:
        ident=_text(v.get("样品编号")); top=num(v.get("顶深_m"),c.get("顶深_m",""),True); bot=num(v.get("底深_m"),c.get("底深_m",""),True); core=num(v.get("岩心长_m"),c.get("岩心长_m",""))
        if not ident: issues.append(_issue("error","MISSING_ID","样品编号不能为空",[c.get("样品编号","")]))
        elif ident in seen["样品"]: issues.append(_issue("error","DUPLICATE_ID",f"样品编号重复：{ident}",[c["样品编号"]]))
        seen["样品"].add(ident)
        if top is not None and bot is not None and (top<0 or bot>endpoint if endpoint is not None else False): issues.append(_issue("error","SAMPLE_OUT_OF_RANGE","样品区间超出孔深范围",[c.get("顶深_m",""),c.get("底深_m","")]))
        if top is not None and bot is not None and bot<=top: issues.append(_issue("error","INVALID_INTERVAL","样品底深必须大于顶深",[c.get("顶深_m",""),c.get("底深_m","")]))
        if core is not None and top is not None and bot is not None and core>bot-top+1e-6: issues.append(_issue("error","CORE_LENGTH_EXCEEDS_INTERVAL","岩心长不能大于样品段长",[c.get("岩心长_m","")]))
        if core is not None and core<0: issues.append(_issue("error","NEGATIVE_CORE_LENGTH","岩心长不能为负数",[c.get("岩心长_m","")]))
        assays={}
        for e in units:
            raw_assay=v.get(e)
            if not _is_blank(raw_assay) and isinstance(raw_assay,str) and any(token in raw_assay for token in ('<','检出','低于')):
                issues.append(_issue("error","NON_NUMERIC_ASSAY_TEXT",f"{e}分析结果为原始文本“{raw_assay}”，保留原值且不转为0；当前图形要求数值",[c.get(e,"")]))
            val=num(raw_assay,c.get(e,""))
            if val is not None:
                if val<0: issues.append(_issue("error","NEGATIVE_ASSAY",f"{e}分析值不能为负数",[c.get(e,"")]))
                if not units[e]: issues.append(_issue("error","MISSING_ASSAY_UNIT",f"填写了{e}分析值但未提供单位",[c.get(e,""),pcells.get(f"{e}单位","项目!B:B")]))
            assays[e]=val
        samples.append({"id":ident,"top_m":top,"bottom_m":bot,"length_m":bot-top if top is not None and bot is not None else None,"core_m":core,"recovery_percent":core/(bot-top)*100 if core is not None and top is not None and bot is not None and bot>top else None,"assays":assays,"source":_source("样品",rn,c)})
    for rn,v,c in recordsets["孔径"]:
        dep=num(v.get("孔深_m"),c.get("孔深_m",""),True); dia=num(v.get("孔径_mm"),c.get("孔径_mm",""),True)
        if dia is not None and dia<=0: issues.append(_issue("error","INVALID_DIAMETER","孔径必须大于0",[c.get("孔径_mm","")]))
        if dep is not None and endpoint is not None and not 0<=dep<=endpoint: issues.append(_issue("error","STRUCTURE_OUT_OF_RANGE","孔径记录深度超出孔深范围",[c.get("孔深_m","")]))
        structures.append({"depth_m":dep,"diameter_mm":dia,"source":_source("孔径",rn,c)})
    # Overlap is permitted (e.g. duplicate sampling); make it explicit and render in separate tracks.
    ordered_samples=sorted((s for s in samples if s["top_m"] is not None and s["bottom_m"] is not None),key=lambda s:s["top_m"])
    farthest=None
    for sample in ordered_samples:
        if farthest and sample["top_m"] < farthest["bottom_m"]-1e-9:
            issues.append(_issue("warning","OVERLAPPING_SAMPLES",f"样品区间存在重叠，将分轨显示：{farthest['id']} / {sample['id']}",[farthest["source"]["cells"].get("顶深_m",""),sample["source"]["cells"].get("顶深_m","")]))
            break
        if farthest is None or sample["bottom_m"]>farthest["bottom_m"]: farthest=sample
    wb.close()
    pattern_path = Path(__file__).resolve().parent / "templates" / "drill-patterns.json"
    pattern_cfg = __import__("json").loads(pattern_path.read_text(encoding="utf-8"))
    code_to_pattern = pattern_cfg["codes"]
    material_to_code = pattern_cfg["materials"]
    for layer in layers:
        code = layer["material_code"]
        expected = material_to_code.get(layer["lithology_name"])
        cell = layer["source"]["cells"].get("花纹代码", "")
        if code and code not in code_to_pattern:
            issues.append(_issue("error", "UNKNOWN_PATTERN_CODE", f"未配置花纹代码：{code}", [cell]))
        elif code and expected and code != expected:
            issues.append(_issue("error", "PATTERN_MATERIAL_CONFLICT", "花纹代码与精确岩性名称映射冲突", [cell, layer["source"]["cells"].get("岩性名称", "")]))
        elif code:
            layer["pattern_id"] = code_to_pattern[code]
        elif expected:
            layer["material_code"] = expected
            layer["pattern_id"] = code_to_pattern[expected]
        else:
            issues.append(_issue("warning", "PENDING_PATTERN", f"岩性“{layer['lithology_name'] or '未提供'}”无精确花纹映射，保留留白待配置", [layer["source"]["cells"].get("岩性名称", "")]))
    if any(i["severity"]=="error" for i in issues): raise GeometryInputError(issues)
    samples.sort(key=lambda s:(s["top_m"],s["bottom_m"],s["id"]))
    source={"filename":path.name,"sha256":_hash(path),"adapter":"drill-1.0"}
    pending=sum(1 for z in layers if not z.get("material_code"))
    return {"drawing_type":"drill","schema_version":"drill-1.0","source":source,"project":{"name":_text(params.get("项目名称")),"hole_id":hole_id,"analysis_units":units,"source_cells":pcells},"meta":{"hole_id":hole_id,"endpoint_m":endpoint,"depth_basis":basis},"turns":turns,"layers":layers,"samples":samples,"structures":structures,"issues":issues,"summary":{"turns":len(turns),"layers":len(layers),"samples":len(samples),"endpoint_m":endpoint,"pending_patterns":pending}}

def load_drill_workbook(path: str | Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("读取钻孔 XLSX 需要 openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=False)
    integrated = ("钻孔基本信息" in wb.sheetnames or
                  {"自定义测试项目", "样品测试结果"}.issubset(wb.sheetnames))
    if integrated:
        wb.close()
        return load_integrated_drill_workbook(path)
    try:
        return _load_drill_workbook_open(path, wb)
    finally:
        wb.close()
